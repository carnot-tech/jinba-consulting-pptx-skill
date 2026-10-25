#!/usr/bin/env python3
"""スライド規約の機械チェック（PPTX / HTML 共通）。

使い方:
  python3 check_deck.py out.pptx
  python3 check_deck.py deck.html
  python3 check_deck.py assets/SuperTemplate_36type.pptx --template   # テンプレ集そのものを検査するとき
  python3 check_deck.py deck.html --forbid ~/.config/deck-forbidden-terms.txt   # 顧客名・社内語の残りを FAIL
  python3 check_deck.py pages.pptx --house house.skin.json   # 既存の資料へ差し込むページ（slide-rules §8.7）
  python3 check_deck.py out.pptx --xml-only                  # PowerPoint が開けなくなる不正だけを見る
  python3 check_deck.py deck.html --old-colors local/old-colors.txt   # 自社のレガシー色を FAIL（local/old-colors.txt は指定なしでも読む）
  python3 check_deck.py deck.html --brand none                # レガシー色の検査だけを外す（他社の書式で組むとき）

正典: references/slide-rules.md
終了コード: FAIL があれば 1。
"""
import re
from html.parser import HTMLParser
import sys
import zipfile
from pathlib import Path

# 自ブランドで禁止するレガシー色（HEX 6桁・#なし）。本体は空のまま、local/old-colors.txt（1行1色）か
# --old-colors で渡す。--brand none で外す（他社の書式で組むとき）
OLD_COLORS = []
LOCAL_DIR = Path(__file__).resolve().parent.parent / "local"
BRANDS = ("none",)  # --brand で受け付ける値。none = レガシー色の検査を外す
TITLE_MAX = 40
EMU_W, EMU_H = 12192000, 6858000

fails, warns = [], []


def fail(msg):
    fails.append(msg)


def warn(msg):
    warns.append(msg)


# --- PowerPoint が開けなくなる不正（slide-rules §8.7）-------------------------------
# テーマ色の欄に HEX や存在しない名前を書いた PPTX は、python-pptx では読めても PowerPoint は「修復」を求める。
# 確認のために自動で開かせると、そのダイアログが残って以後の操作を止める。開く前にここで止める。
SCHEME_COLORS = {"bg1", "tx1", "bg2", "tx2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6",
                 "hlink", "folHlink", "dk1", "lt1", "dk2", "lt2", "phClr"}
XML_ONLY = False   # --xml-only: 上の不正だけを見て終える（PowerPoint に開かせる前に使う）
HOUSE = None       # --house: measure_deck.py が書き出した、差し込む先の資料の書式


def check_xml_validity(idx, xml):
    bad = sorted(set(re.findall(r'<a:schemeClr val="([^"]*)"', xml)) - SCHEME_COLORS)
    if bad:
        fail(f"p{idx}: テーマ色の名前が不正 {bad}（HEX は srgbClr に書く。このままでは PowerPoint が開けない — §8.7）")
    if re.search(r'<a:ext cx="-\d+"|<a:ext cx="\d+" cy="-\d+"', xml):
        fail(f"p{idx}: 図形の大きさが負の値（PowerPoint が開けない — §8.7）")
    szs = [int(v) for v in re.findall(r'<a:(?:rPr|endParaRPr|defRPr)[^>]*\bsz="(\d+)"', xml)]
    if any(not 100 <= v <= 400000 for v in szs):
        fail(f"p{idx}: 文字の大きさ sz が範囲外（100〜400000。PowerPoint が開けない — §8.7）")


def check_dangling_refs(idx, xml, rels_xml):
    """図形・画像・動画を消したあとに残った参照。PowerPoint で開くと修復ダイアログが出る（LibreOffice の PDF 化では気づけない）。"""
    ids = set(re.findall(r'Id="([^"]+)"', rels_xml or ""))
    used = set(re.findall(r'r:(?:embed|id|link|pict)="([^"]+)"', xml))
    missing = sorted(used - ids)
    if missing:
        fail(f"p{idx}: 関係ファイル（rels）に無い参照 {missing[:4]}（消した図形・画像の参照が残っている。PowerPoint が修復を求める — §8.7）")
    timing = re.search(r"<p:timing>.*?</p:timing>", xml, re.S)
    if timing:
        shape_ids = set(re.findall(r'<p:cNvPr[^>]*\bid="(\d+)"', xml))
        orphan = sorted(set(re.findall(r'spid="(\d+)"', timing.group(0))) - shape_ids, key=int)
        if orphan:
            fail(f"p{idx}: アニメーションが存在しない図形を指している（spid {orphan[:4]}。図形を消したらアニメーションも消す — §8.7）")


def check_house(idx, slide, xml):
    """差し込む先の資料と作りが合っているか（slide-rules §8.7）。HOUSE は measure_deck.py の出力"""
    if HOUSE.get("title") and slide.shapes.title is None:
        fail(f"p{idx}: タイトルがプレースホルダーに入っていない（位置と書体は資料のレイアウトが決める — §8.7）")
    faces = sorted(set(re.findall(r'<a:(?:latin|ea) typeface="((?!\+)[^"]*)"', xml)))
    if faces and HOUSE["fonts"]["inherit"]:
        fail(f"p{idx}: 書体を run に直指定 {faces[:4]}（資料はテーマ任せ。差し込むと書体がずれる — §8.7）")
    elif faces and set(faces) - {HOUSE["fonts"].get("explicit")}:
        warn(f"p{idx}: 資料と違う書体 {sorted(set(faces) - {HOUSE['fonts'].get('explicit')})[:4]}（資料は {HOUSE['fonts'].get('explicit')}）")
    if "<a:tbl>" in xml and not HOUSE["tables"]["native"]:
        warn(f"p{idx}: PowerPoint の表オブジェクト（資料の表はテキストボックス＋罫線。deck_pptx の table で組む — §8.7）")
    if 'type="slidenum"' in xml and HOUSE.get("page_number") == "layout":
        warn(f"p{idx}: ページ番号をスライドに置いている（資料はレイアウトが出すので二重になる — §8.7）")
    used = set(v.upper() for v in re.findall(r'<a:srgbClr val="([0-9A-Fa-f]{6})"', xml)) | set(re.findall(r'<a:schemeClr val="([^"]+)"', xml))
    ng = (HOUSE.get("colors") or {}).get("ng") or "FF0000"
    extra = sorted(used - set(HOUSE.get("palette") or used) - {ng.split("|")[0].upper(), "FF0000", "FFFFFF", "000000", "bg1", "tx1"})
    if extra:
        warn(f"p{idx}: 資料に無い色 {extra[:6]}（資料のテーマ色から選ぶ — §8.7）")


# テンプレ集そのものを検査するときだけ True（--template）。
# 通常のデッキでは「Text N」「◯◯」等のプレースホルダーが残っていたら FAIL にする。
TEMPLATE_MODE = False
STRICT_LEN = True  # 互換のため残置。PPTX/HTMLとも2行許容（>80字のみFAIL）

# 表記ゆれ代表ペア（slide-rules §7.6: 1資料1用語）。両方の表記が同一資料に現れたら WARN。
# (ラベルA, パターンA, ラベルB, パターンB)
PART_TYPE_NAMES = set()

TERM_VARIANTS = [
    ("メモリー", r"メモリー", "メモリ", r"メモリ(?!ー)(?!・CPU)"),  # 計算機の「メモリ・CPU」は除外
    ("紐付", r"紐付", "紐づ", r"紐づ"),
    ("ユーザー", r"ユーザー", "利用者", r"利用者"),
    ("アセット", r"アセット", "資産", r"資産"),
    ("フォルダー", r"フォルダー", "フォルダ", r"フォルダ(?!ー)"),
    ("サーバー", r"サーバー", "サーバ", r"サーバ(?!ー)"),
    ("メンバー", r"メンバー", "メンバ", r"メンバ(?!ー)"),
    ("コンピューター", r"コンピューター", "コンピュータ", r"コンピュータ(?!ー)"),
    ("問い合わせ", r"問い合わせ", "問合せ", r"問合せ"),
]


# AI臭ワード（slide-rules §7.9 / references/ai-smell-lexicon.md）。高確度語のみ WARN。
AI_SMELL_WORDS = [
    "まさに", "非常に", "極めて", "圧倒的", "画期的", "革新的", "次世代の",
    "過言ではありません", "に他なりません", "シームレス", "シナジー", "ソリューション",
    "エンドツーエンド", "ブラッシュアップ", "付加価値",
    "寄り添い", "伴走し", "二人三脚", "さらなる高みへ", "邁進",
    "昨今", "変化の激しい", "という点において", "の観点から",
    "させていただきます", "いただけますと幸いです",
    "と言えるでしょう", "と考えられます", "することが可能です",
    # 先送りの決まり文句（中身を書かずに後へ回す）
    "詳細は別途", "別途ご説明", "追ってご連絡", "今後検討してまいります", "詳細は後日",
]


def check_production_meta(slides):
    """制作メタ（本スキル名・リポジトリ）がスライドに見えていないか（SKILL.md「本スキル使用の注釈」）。
    ツール名を置けるのは最終ページ（裏表紙）の出典行だけ。それ以外のページに出ていたら FAIL。
    「本資料は…で作成」の文は免責文のこともあるので、最終ページ以外にあれば WARN で目視に回す。"""
    def visible(s):
        s = re.sub(r"<(script|style)[^>]*>.*?</\1>|<!--.*?-->", " ", s, flags=re.S)
        return re.sub(r"<[^>]+>", " ", s)
    if TEMPLATE_MODE:   # パーツ集は見本の並びなので裏表紙が最終ページにない
        return
    for i, s in enumerate(slides, 1):
        if i == len(slides):
            continue
        v = visible(s)
        m = re.search(r"consulting-pptx-skill", v)
        if m:
            fail(f"p{i}: 制作メタ情報がスライドに表示されている: 「{m.group(0)}」（ツール名を置けるのは最終ページの出典行だけ）")
            continue
        m = re.search(r"本資料は[^。<]{0,40}で作成", v)
        if m:
            warn(f"p{i}: 「{m.group(0)}」（制作クレジットなら最終ページの出典行へ。免責文なら可）")
        elif re.search(r"github\.com", v):
            warn(f"p{i}: スライドに github.com が表示されている（制作メタなら消す。出典として公開リポジトリを示すなら可）")


def pptx_text(xml):
    """スライド XML の表示テキストを段落（<a:p>）ごとに 1 行にして返す。
    run（<a:t>）ごとに改行すると「ラベル」「 — 」「説明」のように run が分かれたダッシュ連結を見落とす。"""
    paras = re.findall(r"<a:p\b.*?</a:p>", xml, re.S)
    if not paras:
        return "\n".join(re.findall(r"<a:t>(.*?)</a:t>", xml, re.S))
    return "\n".join("".join(re.findall(r"<a:t>(.*?)</a:t>", p, re.S)) for p in paras)


def check_ai_smell(pages):
    """slide-rules §7.9: AI臭の高確度語を検出（WARN。文脈上正当なら目視で無視してよい）"""
    hits = {}
    for i, txt in pages:
        found = [w for w in AI_SMELL_WORDS if w in txt]
        if found:
            hits[i] = found
    if hits:
        detail = "、".join(f"p{i}:「{'/'.join(ws[:3])}」" for i, ws in sorted(hits.items())[:6])
        warn(f"AI臭ワード検出: {detail}（§7.9 / ai-smell-lexicon.md。素の動詞・直球の言い方に置き換え）")
    # 非該当セルの「—」単独（§6）は連結ではないので除外する。行（文章塊）単位で、
    # 「—」が他の文字と同じ塊に入っているものだけを数える（「ラベル — 説明」「A—B」）。
    def _dash_joined(txt):
        for line in txt.split("\n"):
            t = line.strip()
            if "—" not in t or re.fullmatch(r"—+(?:\s*[（(][^）)]*[）)])?", t):
                continue
            return True
        return False
    dash_pages = sorted({i for i, txt in pages if _dash_joined(txt)})
    if dash_pages:
        warn(f"ダッシュ「 — 」連結 p{dash_pages}（AI文体の典型。句点・「：」・括弧に置き換え。§7.9）")


def check_terms(pages):
    """pages: [(idx, text), ...] 資料全体で両方の表記が出たら WARN（意味が別なら目視で無視してよい）"""
    for la, pa, lb, pb in TERM_VARIANTS:
        hits_a = sorted({i for i, t in pages if re.search(pa, t)})
        hits_b = sorted({i for i, t in pages if re.search(pb, t)})
        if hits_a and hits_b:
            warn(f"表記ゆれ疑い: 「{la}」p{hits_a} と「{lb}」p{hits_b} が混在（§7.6 1資料1用語。別概念なら可・目視確認）")

# --- 数値の平仄（slide-rules §7.6）--------------------------------------------------
# 同じ指標（数の直前の語）が、別のページで違う値で書かれていたら WARN。
# 例:「売上高 120億円」(p3) と「売上高は約118億円」(p7)、「利用者数 4.2万人」と「利用者数 42,500人」。
# 単位の桁（千・万・億・兆）は掛けて比べるので、書き方が違うだけで値が同じなら出さない。
# 時点（2024年／2024年度）が直前にあれば指標名に含めて比べるので、年の違う同じ指標は食い違いにしない。
NUM_UNITS = (r"円|ドル|ユーロ|人|名|件|社|団体|店舗|拠点|校|戸|世帯|台|個|本|枚|冊|回|倍|%|％|pt|ポイント|"
             r"時間|日|か月|ヶ月|週|kWh|MWh|GWh|kW|MW|GW|kg|km|m3|㎥|m2|㎡|ha|ヘクタール|g|t|トン|m|L")
NUM_LABEL = re.compile(
    r"(?:(?P<year>(?:19|20)\d{2})年度?(?:の|時点の|末の|末時点の)?)?"
    r"(?P<label>[一-龥々ァ-ヶー]{2,12})(?:は|が|の|：|:|＝|=)?\s*(?:約|およそ|計|合計)?\s*"
    r"(?P<num>\d[\d,]*(?:\.\d+)?)\s*(?P<unit>万|億|兆|千)?(?P<base>" + NUM_UNITS + r")(?![A-Za-z])")
_ZEN = str.maketrans("０１２３４５６７８９，．％：＝", "0123456789,.%:=")
_MULT = {"千": 1e3, "万": 1e4, "億": 1e8, "兆": 1e12}
_UNIT_ALIAS = {"名": "人", "％": "%", "ポイント": "pt", "ヶ月": "か月", "㎥": "m3", "㎡": "m2",
               "ヘクタール": "ha", "トン": "t"}
# 指標名にならない一般語（「以上 80%」「平均2.1%」など）は比べない
_GENERIC_LABELS = {"以上", "以下", "未満", "超", "平均", "合計", "全体", "最大", "最小", "最高", "最低",
                   "前年", "前年比", "同期", "うち", "残り", "目標", "実績", "約", "計"}


def _num_facts(text):
    """(指標名, 単位) -> 値 の組を拾う。値は単位の桁（万・億）を掛けた数で比べる。"""
    out = []
    for m in NUM_LABEL.finditer(text.translate(_ZEN)):
        if m["label"] in _GENERIC_LABELS:
            continue
        label = (m["year"] + "年:" if m["year"] else "") + m["label"]
        base = _UNIT_ALIAS.get(m["base"], m["base"])
        try:
            val = round(float(m["num"].replace(",", "")) * _MULT.get(m["unit"] or "", 1), 6)   # 1.1億 = 110,000,000 を float 誤差で食い違いにしない
        except ValueError:
            continue
        out.append(((label, base), val, m.group(0).strip()))
    return out


def check_number_consistency(pages):
    """pages: [(idx, text), ...] 同じ指標がページ間で違う値なら WARN（別概念・別時点なら目視で無視してよい）"""
    seen = {}
    for i, t in pages:
        for key, val, raw in _num_facts(t):
            seen.setdefault(key, []).append((i, val, raw))
    for (label, base), hits in seen.items():
        vals = {v for _, v, _ in hits}
        pages_hit = {i for i, _, _ in hits}
        if len(vals) > 1 and len(pages_hit) > 1:
            detail = " / ".join(f"p{i}「{raw}」" for i, _, raw in hits[:4])
            warn(f"数値の平仄疑い: 「{label.replace(':', '')}」が {detail} で食い違う（§7.6 数値の平仄。別時点・別範囲なら注記して可）")

# --- 本文のプレースホルダー残り（slide-rules §2.8 / README）-------------------------
# タイトルだけでなく本文・表・カードに「Text N」「ラベル N」「YYYY」「パーツNN｜」が残っていたら FAIL。
# 型名をそのままタイトルにしたページ（例:「軸のある表」）も FAIL。
BODY_PLACEHOLDER = re.compile(
    r"Text\s*\d+|ラベル\s*\d+|タイトル\s*\d+|Source\s*\d+|YYYY|XXX|ＸＸＸ|○○|〇〇|●●|△△|\bTBD\b|Lorem|パーツ\s*\d+\s*[｜|]|ダミー|^会社名$|^連絡先$")


def _part_type_names():
    """パーツ集の型名（new_deck.py --list と同じ並び）。テンプレが無い環境では空。"""
    names = set()
    root = Path(__file__).resolve().parent.parent / "templates"
    for f in ("freeform_parts_16x9.html", "freeform_parts_more_16x9.html"):
        fp = root / f
        if not fp.exists():
            continue
        for m in re.finditer(r"パーツ\s*\d+\s*[｜|]\s*([^<]+)<", fp.read_text(encoding="utf8", errors="ignore")):
            names.add(m.group(1).strip())
    return names


def check_body_placeholders(idx, leaf_texts, title):
    if TEMPLATE_MODE:
        return
    hits = sorted({t for t in leaf_texts if BODY_PLACEHOLDER.search(t)})
    if hits:
        fail(f"p{idx}: 本文にテンプレのプレースホルダーが残っている ×{len(hits)}: {' / '.join(h[:20] for h in hits[:4])}")
    if title and title.strip() in PART_TYPE_NAMES:
        fail(f"p{idx}: タイトルがパーツの型名のまま「{title.strip()}」（主張文に書き換える — §2.8）")


# --- 1ブロック1文（文章を連続して詰め込まない）------------------------------------
# スライドの1つの文章塊（セル・カード本文・段落・箇条1行）に2文以上を入れない。
# 2文以上になる中身は1項目1文の箇条に分ける。機械判定は「句点で終わる文が2つ以上」。
SENTENCE_END = re.compile(r"[。！？!?](?=\s*\S)")


def check_multi_sentence(idx, leaf_texts):
    bad = []
    for t in leaf_texts:
        if t.startswith(("出典", "注", "※", "Source")):
            continue  # 出典・注記行は対象外
        core = re.sub(r"（[^）]*）|\([^)]*\)|「[^」]*」", "", t)  # 括弧・引用内の句点は数えない
        if len(SENTENCE_END.findall(core.strip())) >= 1:
            bad.append(t)
    if bad:
        fail(f"p{idx}: 1つの文章塊に2文以上 ×{len(bad)}: 「{bad[0][:36]}…」（1項目1文の箇条に分ける。動きはきっかけ→起きること→判断→結果の順）")


# --- 1文の項目の文末「。」（§7.25）---------------------------------------------------
def check_trailing_period(idx, leaf_texts):
    """1文で終わる箇条・ラベル・表のマスの文末に「。」を付けない。出典・注記は対象外。"""
    bad = [t for t in leaf_texts
           if t.endswith("。") and not t.startswith(("出典", "注", "※", "Source"))
           and not SENTENCE_END.search(t[:-1])]
    if bad:
        warn(f"p{idx}: 1文の項目の文末に「。」 ×{len(bad)}: 「{bad[0][-24:]}」（§7.25: 文末の句点を外す）")


# --- 本文に直接打ったブレット記号（§7.3）---------------------------------------------
TYPED_BULLET = re.compile(r">\s*([・●•◆■▪◦‣]|-\s|－)\s*[^\s<]")


def check_typed_bullets(idx, fragment):
    """同じページでブレット記号の形を揃える。本文に打った記号が2種類以上、または打った記号と li のリストが混ざると WARN。"""
    body = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", fragment, flags=re.S)
    kinds = sorted(set(h.strip() for h in TYPED_BULLET.findall(body)))
    if len(kinds) >= 2:
        warn(f"p{idx}: 本文に打ったブレット記号が混ざっている（{''.join(kinds)}）（§7.3: 1つのデッキで記号の形を揃える。リストは li::before を既定にする）")
    elif kinds and re.search(r"<li\b", body):
        warn(f"p{idx}: 本文に打った記号（{kinds[0]}）と li のリストが同じページに混ざっている（§7.3: どちらかに揃える）")


# --- 同じ見せ方のページの連続（§4.96）------------------------------------------------
def page_shape(fragment):
    if re.search(r"<table\b", fragment):
        return "表"
    return ""


def check_shape_runs(shapes):
    """表のページが3枚以上続いたら WARN。意図して並べた連番（同じ型の比較を続けて見せる等）なら無視してよい。"""
    if TEMPLATE_MODE:
        return  # パーツ集は同じ型を並べて見せる見本なので対象外
    run = []
    for i, sh in enumerate(shapes + [""], 1):
        if sh and run and shapes[run[-1] - 1] == sh:
            run.append(i)
            continue
        if len(run) >= 3:
            warn(f"p{run[0]}〜p{run[-1]}: {shapes[run[0]-1]}のページが{len(run)}枚続く（§4.96: 比較・流れ・因果・位置づけのどれを見せるかを問い直し、形を変える。同じ型の比較を続けて見せるのが目的のときだけ残し、理由を qa.md に書く）")
        run = [i] if sh else []


def check_page_refs(pages):
    """ページを足す・消す・並べ替えたあとに残った参照（「→ P.n」・目次の番号）を拾う。
    ページ数を超える番号を指していたら WARN（範囲内の取り違えは目視。§8 チェックリスト）。"""
    if TEMPLATE_MODE or not pages:
        return
    total = len(pages)
    for idx, text in pages:
        bad = sorted({int(n) for n in re.findall(r"(?<![A-Za-z])[PpＰ]\s*[.．]\s*([0-9]{1,3})(?![0-9])", text)
                      if int(n) > total})
        if bad:
            warn(f"p{idx}: 「P.{bad[0]}」はページ数（{total}枚）を超えている（構成を変えたあとの参照・目次の番号を振り直す）")


# --- 配布物に残った操作用の部品（§1）------------------------------------------------
def check_static_ui(html):
    """PDF・印刷・静的な配布物に、押せるように見えるボタン・タブ・ページ送りを残さない。印刷時に隠していれば可。"""
    body = html.split("</style>", 1)[1] if "</style>" in html else html
    n = len(re.findall(r"<button\b|role=\"tab\"|class=\"[^\"]*\b(?:pager|nav-btn|prev|next|tabbar)\b", body))
    if n and not re.search(r"@media\s+print\s*{[^@]*(button|\.pager|\.nav|\[role=.?tab)[^{]*{[^}]*display\s*:\s*none", html, re.S):
        warn(f"操作用の部品（ボタン・タブ・ページ送り）×{n} が印刷時に隠れていない（§1: @media print で display:none にする）")


def check_mixed_parts(html, slide_classes):
    """基本パーツ集と追加パーツ集（1600px 基準）を混ぜたデッキ。追加パーツ側の縮小が抜けるとそのページだけ大きく出る。"""
    if re.search(r"--slide-w\s*:\s*1600px", html) and re.search(r"338\.67mm", html) and len(set(slide_classes)) > 1:
        if not re.search(r"\.slide\s*{[^}]*zoom\s*:\s*0?\.8", html):
            warn("基本パーツ集と追加パーツ集が混在しているが、追加パーツの .slide に zoom:0.8 が無い（§4.97: ページごとに文字の大きさがずれる。check_layout.mjs で大きさの揃いを確かめる）")


# --- 禁止語（顧客名・社内語・案件コード）-------------------------------------------
# 公開・社外共有前に、資料ごとに出してはいけない語を外部ファイルで渡す（リポジトリには入れない）。
#   python3 check_deck.py deck.html --forbid ~/.config/deck-forbidden-terms.txt
# 1行1語。# で始まる行はコメント。re: で始めると正規表現。
FORBIDDEN_TERMS = []


def load_forbidden(path):
    terms = []
    for line in Path(path).expanduser().read_text(encoding="utf8").splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        terms.append(re.compile(line[3:]) if line.startswith("re:") else re.compile(re.escape(line)))
    return terms


def check_forbidden(pages, raw=""):
    if not FORBIDDEN_TERMS:
        return
    for pat in FORBIDDEN_TERMS:
        hit_pages = sorted({i for i, t in pages if pat.search(t)})
        in_markup = bool(raw and pat.search(re.sub(r">[^<]*<", "><", raw)))  # コメント・属性・alt 等
        if hit_pages or in_markup:
            where = f"p{hit_pages}" if hit_pages else "HTMLのコメント/属性"
            # 語そのものは出力に出さない（ログ経由の再流出を防ぐ）
            fail(f"禁止語リストの語が {where} に残っている（リスト{FORBIDDEN_TERMS.index(pat)+1}行目の語）")


CLOSING_TITLE = re.compile(r"^(ご清聴)?ありがとうございました[。！!]?$|^Thank you", re.I)


def check_title(idx, title, explicit_break=False, closing=False, edge=False):
    # explicit_break: 意味の切れ目で明示改行済み（§2.13）なら2行想定の WARN は出さない
    # closing: 表紙・裏表紙・締めのあいさつだけのページ。語尾の検査から外す
    # edge: class に cover/chap/back が付いたページ。タイトルが無くてよい
    t = title.strip()
    if not t:
        if not edge:
            warn(f"p{idx}: タイトルが空（表紙/扉なら可）")
        return
    if re.search(r"(\d+|[一二三四五六七八九十]+)(つ|点|観点|方向|段階|要素|領域|論点|問い)(に|で|へ)?(を)?(整理|分解|分け|まとめ|分類|束ね)(する|る|た)?$|順に見る$", t):
        warn(f"p{idx}: タイトルが作業の説明になっている（「〜に分ける」「〜で整理する」は作り手の手順。分けた結果、何が言えるかを書く — §2.12・§2.19）: 「{t}」")
    elif re.search(r"を(一覧に|一覧化)(する|した)$", t):
        fail(f"p{idx}: タイトルが作業の説明で終わっている（「〜を一覧にする」は作り手の手順。並べた結果、何が言えるかを主張で書く — §2.12・§2.19）: 「{t}」")
    elif re.search(r"(ごとに|別に)、?([^、・をの]+・)+[^、・をの]+を(整理(する|した)|まとめる|まとめた)$", t):
        # 「打ち手ごとに担当・時期・効果を整理する」: 切り口（〜ごとに）と項目の並び（A・B・C）を「整理する」＝資料の作業の説明
        # 中黒だけ・「ごとに」だけでは FAIL にしない（「申請・承認の重複した窓口をまとめる」は打ち手の言い切り）
        fail(f"p{idx}: タイトルが作業の説明で終わっている（「〜ごとに…を整理する」は作り手の手順。整理した結果、何が言えるかを主張で書く — §2.12・§2.19）: 「{t}」")
    elif re.search(r"を整理(する|した)$|を(まとめる|まとめた)$", t):
        # 「承認フローを整理する」「重複した窓口をまとめる」のように打ち手の言い切りでもあり得るので WARN（上の「3つに整理する」と揃える）
        warn(f"p{idx}: タイトルが「〜を整理する」「〜をまとめる」で終わっている（資料の作業の説明なら、整理した結果、何が言えるかを書く。打ち手として何をどうするかを言い切っているなら残してよい — §2.12・§2.19）: 「{t}」")
    elif re.search(r"を(示す|説明する|紹介する)$|について(述べる|説明する)$", t):
        warn(f"p{idx}: タイトルが作業の説明で終わっている（「〜を示す」は何を見せるかの説明。見せた結果、何が言えるかを書く — §2.12）: 「{t}」")
    if not closing and re.search(r"(です|ます|でした|ました)[。．.]?$", t):
        fail(f"p{idx}: タイトルがですます調 → 体言止めに: 「{t}」")
    # タイトルは PPTX/HTML とも2行まで許容。1行目安(TITLE_MAX=40字)超は WARN、2行にも収まらない長さ(>80字)だけ FAIL。
    # 文字を縮小して1行に詰めるのは不可（slide-rules §2.1）。
    # 半角文字（英数・記号・空白）は全角の半分として数える（英語タイトルを日本語基準で FAIL にしない）
    tlen = sum(0.5 if ord(ch) < 0x3000 else 1 for ch in t)
    if tlen > TITLE_MAX * 2:
        fail(f"p{idx}: タイトル 全角換算{tlen:.0f}字（>{TITLE_MAX*2}・2行にも収まらない。主張を絞る）: 「{t}」")
    elif tlen > TITLE_MAX and not explicit_break:
        warn(f"p{idx}: タイトル 全角換算{tlen:.0f}字（2行になる想定。意味の切れ目で改行・泣き別れなし・文字縮小で1行に詰めない）: 「{t}」")
    if re.match(r"^(Step|STEP|ステップ)\s*\d", t):
        fail(f"p{idx}: タイトルに Step 連結（タグチップで表現）: 「{t}」")
    if re.search(r"^(この|その|ここまで)", t):
        fail(f"p{idx}: 他スライド参照語で始まるタイトル: 「{t}」")
    if re.match(r"^(まずは|では|そして|さらに|ちなみに)|^(まず|また|次に)[、,]", t):
        warn(f"p{idx}: タイトルが話し言葉の接続詞で始まる（主語から書く — slide-rules §2.18）: 「{t}」")
    if t.count("（") + t.count("(") >= 2:
        warn(f"p{idx}: タイトルに丸括弧が多い: 「{t}」")
    if not TEMPLATE_MODE and re.search(r"[◯○]{2,}|Text\s*\d|ラベル\s*\d|タイトル\s*\d|Source\s*\d|YYYY|ダミー|^資料名$|^会社名$", t):
        fail(f"p{idx}: テンプレのプレースホルダーが残っている（タイトルはストーリーラインから書く — slide-rules §2.8）: 「{t}」")


# テンプレの見本タイトルは「形の参考」であって埋める鋳型ではない（slide-rules §2.8）。同じ文型が並んだら WARN。
# 「第2段階」のような序数は数で束ねた形ではないので外す
COUNT_GROUP = re.compile(r"(?<![\d.第])([2-9２-９]|[二三四五六])\s*(つの|つに|本の|本柱|点の|点に|領域|つ$|本$|点$)")
# 「3点を承認いただく」「2段階で進める」は数を材料として言っているだけのことも多い。束ねた形の候補として WARN の数にだけ入れる
COUNT_SOFT = re.compile(r"(?<![\d.第])([2-9２-９]|[二三四五六])\s*(点を|段階)")


def check_count_titles(titles):
    """§2.4・§4.103: 数で束ねたタイトル（「3つの壁」「3本柱」「3点の承認」）が3枚以上なら FAIL。型が透けて見える"""
    hits = [t for t in titles if t and COUNT_GROUP.search(t)]
    soft = [t for t in titles if t and t not in hits and COUNT_SOFT.search(t)]
    msg = "（§4.103: 数でなく中身をタイトルに書く。「3本柱で進める」→柱の名前で言う。数は材料で決め、3つに揃えない）"
    if len(hits) >= 3:
        fail(f"数で束ねたタイトルが {len(hits)} 枚: 「{hits[0][:24]}」ほか{msg}")
    elif len(hits) + len(soft) >= 3:
        warn(f"数を前に出したタイトルが {len(hits) + len(soft)} 枚: 「{(hits + soft)[0][:24]}」ほか{msg}")


def check_title_variety(titles):
    real = [t.strip() for t in titles if t and t.strip()]
    if len(real) < 5:
        return
    molds = {
        "「◯◯は、…」": lambda t: re.match(r"^.{1,12}は[、,]", t),
        "「◯◯には、…がある」": lambda t: re.search(r"には[、,].*(ある|存在する)$", t),
        "数を主語にした形（「3つの…」）": lambda t: re.match(r"^[0-9０-９一二三四五六七八九十]+つ", t),
    }
    for name, f in molds.items():
        hits = [t for t in real if f(t)]
        if len(hits) >= max(4, int(len(real) * 0.6)):
            warn(f"タイトルの文型が {name} に偏っている（{len(hits)}/{len(real)}枚）。テンプレの見本文型をなぞらず、主張ごとに自然な文で書く — slide-rules §2.8")


# --- 数の不一致（slide-rules §2.9） -------------------------------------------
# タイトルに書いた数と、本文の連番ラベルの最大値を突き合わせる。
# 例: タイトル「3段階で移行する」に対し本文が「STEP 1〜4」。
# 機械で拾える数少ない「内容の矛盾」なので FAIL にする。
KANSUJI = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5,
           "六": 6, "七": 7, "八": 8, "九": 9, "十": 10}
COUNT_WORDS = r"(?:段階|フェーズ|ステップ|柱|論点)"
ENUM_FAMILY = [
    (r"段階|フェーズ|ステップ", r"(?:フェーズ|ステップ|STEP|Step|PHASE|Phase|段階)"),
    (r"柱", r"(?:柱)"),
    (r"論点", r"(?:論点)"),
]


def _to_int(s):
    s = s.translate(str.maketrans("０１２３４５６７８９", "0123456789"))
    if s.isdigit():
        return int(s)
    return KANSUJI.get(s)


def check_count_match(idx, title, text):
    """タイトルの「N段階」等が本文の連番ラベルと食い違っていたら FAIL"""
    if not title or not text:
        return
    for m in re.finditer(r"([0-9０-９]+|[一二三四五六七八九十])\s*(" + COUNT_WORDS + ")", title):
        n = _to_int(m.group(1))
        if not n or not (2 <= n <= 12):
            continue
        fam = next((lab for pat, lab in ENUM_FAMILY if re.search(pat, m.group(2))), None)
        if not fam:
            continue
        found = [v for v in (_to_int(x.group(1))
                             for x in re.finditer(fam + r"\s*([0-9０-９]+)", text)) if v]
        if found and max(found) != n:
            fail(f"p{idx}: タイトルの「{m.group(0)}」と本文の連番（最大{max(found)}）が食い違う"
                 f"（§2.9 数はページの中身と一致させる）: 「{title}」")


# ---------------------------------------------------------------- PPTX
# ---- スキル有無検証（slide-rules §2.13 泣き別れ／§5.13 版面充填率）
def _fwlen(t):
    return sum(0.5 if ord(ch) < 0x3000 else 1 for ch in t)

def check_orphan(idx, shape, label="タイトル"):
    """箱幅とフォントサイズから1行容量を推定し、最終行が1〜3字だけになる折返し（泣き別れ）を WARN"""
    try:
        if not (shape.has_text_frame and shape.width):
            return
        tf = shape.text_frame
        if tf.word_wrap is False:  # 折り返しなしの文字（html_to_pptx.py の1行テキスト等）は泣き別れしない
            return
        raw = tf.text
        if not raw.strip():
            return
        szs = [r.font.size.pt for p in tf.paragraphs for r in p.runs if r.font.size]
        if not szs:
            return
        pt = max(szs)
        w_in = shape.width / 914400.0
        cap = max(4, int(w_in / (pt / 72.0)))
        for line in raw.replace("\v", "\n").split("\n"):
            L = _fwlen(line.strip())
            if L <= cap:
                continue
            last = L - cap * ((int(-(-L // cap))) - 1)
            if 0 < last <= 3:
                warn(f"p{idx}: {label}が泣き別れの疑い（推定{int(-(-L // cap))}行目が約{last:.0f}字だけ。意味の切れ目で明示改行 — slide-rules §2.13）: 「{line.strip()[:40]}」")
    except Exception:
        return

def check_fill_ratio(idx, slide, title_shape):
    """本文（タイトル下〜出典行上）の縦幅に対する本文要素の占有率。55%未満は WARN（§5.13）"""
    try:
        top_lim = 1.6 * 914400
        bot_lim = 6.8 * 914400
        tops, bots = [], []
        for sh in slide.shapes:
            if sh is title_shape or sh.top is None or sh.height is None:
                continue
            t, b = sh.top, sh.top + sh.height
            if b <= top_lim or t >= bot_lim:
                continue  # キッカー/タイトル/フッター
            # 幅いっぱいの罫線（フッター罫・タイトル罫）は除外
            if sh.height < 914400 * 0.02 and sh.width and sh.width > 914400 * 11:
                continue
            tops.append(max(t, top_lim)); bots.append(min(b, bot_lim))
        if not tops:
            return
        cov = (max(bots) - min(tops)) / (bot_lim - top_lim)
        if cov < 0.55:
            warn(f"p{idx}: 版面充填率 {cov*100:.0f}%（本文が版面の半分未満。情報を足すか型を変える。飾りで埋めない — slide-rules §5.13）")
    except Exception:
        return


def check_pptx(path):
    try:
        from pptx import Presentation
        from pptx.util import Emu
    except ImportError:
        sys.exit("python-pptx が必要: pip3 install python-pptx")
    from pptx_open import open_presentation
    prs = open_presentation(path)
    want_w, want_h = (HOUSE["slide"]["emu"] if HOUSE else (EMU_W, EMU_H))   # 差し込むページは資料の大きさに合わせる（§1）
    if not XML_ONLY and (abs(prs.slide_width - want_w) > 2000 or abs(prs.slide_height - want_h) > 2000):
        fail(f"スライドサイズ {prs.slide_width}x{prs.slide_height} ≠ {'資料' if HOUSE else '16:9'} {want_w}x{want_h}")

    titles = []
    term_pages = []
    with zipfile.ZipFile(path) as z:
        names = [n for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
        for n in sorted(names, key=lambda s: int(re.search(r"slide(\d+)", s).group(1))):
            xml = z.read(n).decode("utf8", "ignore")
            idx = int(re.search(r"slide(\d+)", n).group(1))
            term_pages.append((idx, pptx_text(xml)))
            check_xml_validity(idx, xml)
            rels_name = f"ppt/slides/_rels/slide{idx}.xml.rels"
            check_dangling_refs(idx, xml, z.read(rels_name).decode("utf8", "ignore") if rels_name in z.namelist() else "")
            if re.search(r'<p:sld\b[^>]*\bshow="0"', xml):
                warn(f"p{idx}: 非表示のスライド（使わないならファイルから消す。付録の控えとして意図して残すなら無視してよい — slide-rules §8）")
            if XML_ONLY:
                continue
            # 角丸（高さ 0.4in=365760 EMU 以上の図形のみ）
            rr = 0
            for m in re.finditer(r"<p:sp>.*?</p:sp>", xml, re.S):
                sp = m.group(0)
                if 'prst="roundRect"' in sp:
                    h = re.search(r'<a:ext cx="\d+" cy="(\d+)"', sp)
                    if h and int(h.group(1)) >= 365760:
                        rr += 1
            if rr:
                fail(f"p{idx}: roundRect の大きなボックス ×{rr}（直角 rect に）")
            # 塗りありボックスに枠線（大きな図形のみ）
            fb = 0
            for m in re.finditer(r"<p:sp>.*?</p:sp>", xml, re.S):
                sp = m.group(0)
                spPr = re.search(r"<p:spPr>.*?</p:spPr>", sp, re.S)
                if not spPr:
                    continue
                pr = spPr.group(0)
                h = re.search(r'<a:ext cx="\d+" cy="(\d+)"', pr)
                if not (h and int(h.group(1)) >= 365760):
                    continue
                ln = re.search(r"<a:ln[ >].*?</a:ln>", pr, re.S)
                filled = "<a:solidFill>" in re.sub(r"<a:ln[ >].*?</a:ln>", "", pr, flags=re.S)
                if filled and ln and "<a:solidFill>" in ln.group(0):
                    fb += 1
            if fb:
                warn(f"p{idx}: 塗りあり図形に枠線 ×{fb}（塗りカードは line なし — slide-rules §5.3）")
            for c in OLD_COLORS:
                if re.search(f'val="{c}"', xml, re.I):
                    fail(f"p{idx}: 禁止色 {c}")
            # ブレット記号のテキスト直打ち（slide-rules §7.3: マーカーは buChar 書式で付与）
            lb = len(re.findall(r"<a:t>\s*•", xml))
            if lb:
                fail(f"p{idx}: ブレット「•」を本文テキストに直打ち ×{lb}（buChar/buNone 書式に — slide-rules §7.3）")
            ld = len(re.findall(r"<a:t>\s*[–‐-]\s\s", xml))
            if ld:
                warn(f"p{idx}: 第2階層マーカー「– 」らしきテキスト直打ち ×{ld}（buChar 書式に — slide-rules §7.3）")
        # theme / master も色チェック
        for n in z.namelist():
            if "theme" in n or "slideMaster" in n or "slideLayout" in n:
                xml = z.read(n).decode("utf8", "ignore")
                for c in OLD_COLORS:
                    if re.search(f'val="{c}"', xml, re.I):
                        warn(f"{n}: 禁止色 {c}")

    if XML_ONLY:
        return titles
    if HOUSE:
        with zipfile.ZipFile(path) as z:
            for i, s in enumerate(prs.slides, 1):
                check_house(i, s, z.read(s.part.partname.lstrip("/")).decode("utf8", "ignore"))
    for i, s in enumerate(prs.slides, 1):
        title = ""
        named = next((sh for sh in s.shapes if sh.has_text_frame and sh.name.startswith("Title")), None)
        if s.shapes.title is not None and s.shapes.title.has_text_frame:
            title = s.shapes.title.text_frame.text
        elif named is not None:  # html_to_pptx.py が「Title」と名付けたテキストボックス
            title = named.text_frame.text
        else:
            # タイトルPH が無いビルダー: 上部 y<1.2in の最大フォントテキストをタイトルとみなす
            cands = []
            for sh in s.shapes:
                if sh.has_text_frame and sh.top is not None and sh.top < Emu(1097280):
                    sz = max((r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size), default=0)
                    cands.append((sz, sh.text_frame.text))
            if cands:
                title = max(cands)[1]
        # 泣き別れ推定（タイトル＋表紙の大きな文字）と版面充填率
        title_shape = None
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text == title:
                title_shape = sh
                break
        if title_shape is not None:
            check_orphan(i, title_shape, "タイトル")
        if i == 1:
            for sh in s.shapes:
                if sh is not title_shape and sh.has_text_frame:
                    szs = [r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size]
                    if szs and max(szs) >= 13 and len(sh.text_frame.text) < 80:
                        check_orphan(i, sh, "表紙の文字")
        else:
            check_fill_ratio(i, s, title_shape)
        explicit_break = "\n" in title.strip() and all(_fwlen(l.strip()) <= TITLE_MAX for l in title.split("\n"))
        title = title.replace("\n", " ")
        titles.append(title)
        check_title(i, title, explicit_break)
        check_count_match(i, title, dict(term_pages).get(i, ""))
        check_axis_names_pptx(i, s)
        others = [sh.text_frame.text.strip() for sh in s.shapes
                  if sh.has_text_frame and sh is not title_shape and sh.text_frame.text.strip()]
        others = [x for x in others if not re.fullmatch(r"\d{1,3}", x)]
        # 表・グラフ・グループ（図）、または画像2枚以上（画像1枚はロゴとみなす）。13=画像・6=グループ
        kinds = [_shape_type(sh) for sh in s.shapes]
        rich = (any(getattr(sh, "has_table", False) or getattr(sh, "has_chart", False) for sh in s.shapes)
                or 6 in kinds or kinds.count(13) >= 2)
        APPX_PAGES[i] = (appendix_tag_pptx(s, title_shape, title, prs.slide_width, prs.slide_height),
                         sum(_fwlen(x) for x in others), rich)
        # サブタイトル疑い: タイトル直下 (y 1.2〜1.75in) の細字テキスト1行
        for sh in s.shapes:
            if sh.has_text_frame and sh.top is not None and Emu(1097280) <= sh.top < Emu(1600200):
                txt = sh.text_frame.text.strip()
                if txt and "\n" not in txt and len(txt) < 60 and txt != title:
                    szs = [r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size]
                    if szs and max(szs) <= 12 and not any(r.font.bold for p in sh.text_frame.paragraphs for r in p.runs):
                        warn(f"p{i}: タイトル直下にサブタイトルらしき行: 「{txt}」")
    PAGES[:] = term_pages
    check_terms(term_pages)
    check_number_consistency(term_pages)
    check_page_refs(term_pages)
    check_ai_smell(term_pages)
    return titles


# ---------------------------------------------------------------- HTML
# §7.22 英字大文字の装飾キッカー／§4.49 接続詞で始まる左右カラム見出し
def check_kicker_and_conclusion(html):
    body = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", html, flags=re.S)
    texts = [re.sub(r"\s+", " ", t).strip() for t in re.findall(r">([^<>]{3,60})<", body)]
    kick = [t for t in texts if re.fullmatch(r"(?:\d{2}\s*[・·/|]\s*)?[A-Z][A-Z0-9 &/·・\-]{5,}", t)
            and re.search(r"[A-Z]{3,}\s+[A-Z]{2,}", t)
            and not re.search(r"(?i)confidential|appendix|section|step|page", t)]
    if kick:
        warn(f"英字大文字の装飾キッカー ×{len(kick)}: {' / '.join(sorted(set(kick))[:4])}（§7.22: 日本語デッキでは右上タグチップで話題を示す）")
    # 左右2カラムの見出し（h3/h4/.hd）だけを見る。th・行見出し（.rh）は行軸で通して読めるので対象外（§4.49）
    heads = [re.sub(r"<[^>]+>", "", t).strip() for t in re.findall(r"<(?:h3|h4|div class=\"(?:hd|colhd|colh|col-h)[^\"]*\")[^>]*>(.*?)</", body, re.S)]
    dakara = [t for t in heads if re.match(r"^(だから|なので|つまり)[、:：]?", t)]
    if dakara:
        warn(f"左右カラムの見出しが接続詞で始まる ×{len(dakara)}（§4.49: 2コンテンツの見出しは単独で読める名詞句に）")


def _leaf_texts(fragment):
    """スライド内の文章塊（ブロック要素ごとのテキスト）。<br> と箇条記号「•」は区切りとして扱う。"""
    frag = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", fragment, flags=re.S)
    frag = re.sub(r"<!--.*?-->", " ", frag, flags=re.S)
    frag = re.sub(r"<br\s*/?>", "\n", frag)
    blocks = re.split(r"</?(?:div|p|li|td|th|h[1-6]|section|ul|ol|tr|table)\b[^>]*>", frag)
    out = []
    for b in blocks:
        for piece in re.split(r"\n|•", re.sub(r"<[^>]+>", "", b)):
            piece = re.sub(r"\s+", " ", piece).strip()
            if piece:
                out.append(piece)
    return out

# 文字化けの痕（置換文字・私用領域・制御文字）。ふだんの日本語の本文には出ない
GARBLED = re.compile(r"[�-\x00-\x08\x0B\x0C\x0E-\x1F\x7F]")


def check_garbled(idx, text):
    """編集・並べ替え・変換の途中で文字が化けた痕を FAIL にする（slide-rules §8）"""
    hits = GARBLED.findall(text)
    if hits:
        cps = sorted({f"U+{ord(c):04X}" for c in hits})
        fail(f"p{idx}: 文字化けの痕 ×{len(hits)}（{' '.join(cps[:4])}）（元の文字に打ち直す。編集画面での並べ替え・変換のあとに起きやすい）")


FRAMEWORK_NAME = re.compile(r"(Situation|Complication|Resolution|So What\??|Why So\??|MECE|SCR|Issue|Key Takeaways?)", re.I)
JA = r"[\u3040-\u30ff\u4e00-\u9fff]"
HALF_COMMA = re.compile(JA + r",\s?" + JA)


# 全角ダッシュ・水平線でのつなぎ。罫線の文字（─━）は前後に空白があるときだけ（「├─ 項目」のような階層表記は対象外）
# 空白は _leaf_texts で1つに正規化済み
DASH_LINK = re.compile(r"\S ?[—―]{1,2} ?\S|\S [─━]{1,2} \S")


def check_dash_link(idx, leaves):
    """§7.1: 「 — 」でつながない（AI 臭の典型）。非該当セルの「—」単独は可。
    ブロック要素ごとの文章塊で見るので、タグの組み方や空白の数で判定が変わらない"""
    for text in leaves:
        m = DASH_LINK.search(text)
        if m:
            fail(f"p{idx}: 「—」で語句をつないでいる「{text[max(0, m.start() - 10):m.end() + 10].strip()}」（句点・「：」・括弧で区切る。§7.1）")
            return


def check_framework_and_commas(idx, leaves, vis):
    """§7.32: 英語の思考の枠組みの名前（Situation／Complication／Resolution 等）を見出し・列名に書かない。
    日本語の文の中の半角カンマは読点「、」にする"""
    hit = [x for x in leaves if FRAMEWORK_NAME.fullmatch(x.strip())]
    if hit:
        fail(f"p{idx}: 思考の枠組みの名前「{hit[0]}」が見出し・列名に出ている（作り手の道具。「これまで／課題／対策」のように中身の言葉にする — §7.32）")
    m = HALF_COMMA.search(vis)
    if m:
        warn(f"p{idx}: 日本語の文の中に半角カンマ「{m.group(0)}」（読点「、」にする）")


WORK_FILES = re.compile(r"(?<![A-Za-z0-9_])(sources|qa|REPORT)\.md(?![A-Za-z0-9_])")  # 日本語に直接続く「のsources.md」も拾う


def check_work_files(idx, text):
    """§2.23: 読み手は作り手の控え（sources.md・qa.md）を見ない。デッキの注でそこを根拠として指さない"""
    m = WORK_FILES.search(text)
    if m:
        fail(f"p{idx}: 作り手の控えのファイル名「{m.group(0)}」がデッキに出ている（読み手は見ない。試算の前提と算式・出典は、そのページの注か付録に書く — §2.23）")


class _ListCounter(HTMLParser):
    """ul/ol ごとに直下の li の数を数える（入れ子も別々に数える）"""
    def __init__(self):
        super().__init__()
        self.stack, self.singles = [], []
    def handle_starttag(self, tag, attrs):
        if tag in ("ul", "ol"):
            self.stack.append([0, ""])
        elif tag == "li" and self.stack:
            self.stack[-1][0] += 1
    def handle_data(self, data):
        if self.stack and self.stack[-1][0] == 1 and len(self.stack[-1][1]) < 30:
            self.stack[-1][1] += data.strip()
    def handle_endtag(self, tag):
        if tag in ("ul", "ol") and self.stack:
            n, txt = self.stack.pop()
            if n == 1:
                self.singles.append(txt)


def check_single_item_lists(idx, fragment):
    """§7.3: 箇条は2項目以上。1つしか書けないなら箇条をやめて地の文にする"""
    pc = _ListCounter()
    try:
        pc.feed(fragment)
    except Exception:
        return
    if pc.singles:
        warn(f"p{idx}: 項目が1つだけの箇条書き ×{len(pc.singles)}: 「{pc.singles[0][:24]}」（§7.3: 箇条は2項目以上。1つなら地の文に）")


def check_body_desu(idx, leaves, title):
    """§7.4: 本文（ブレット・表のセル）も常体。です・ます で終わる項目を WARN（依頼のページの文は目視で可）"""
    bad = [x for x in leaves if x != title and re.search(r"(です|ます|でした|ました)[。．.]?$", x)]
    if bad:
        warn(f"p{idx}: 本文がですます調 ×{len(bad)}: 「{bad[0][-24:]}」（§7.4: 本文も常体で。相手への依頼文だけは可）")


CIRCLED = "①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳"


def check_numbered_labels(pages):
    """§7.13・§7.33: 同じ丸数字の項目名がページによって違う（別の意味に同じ番号／言い換え）"""
    seen = {}
    for idx, text in pages:
        for m in re.finditer(rf"([{CIRCLED}])\s*([^\s{CIRCLED}、。：:（(|｜→]{{2,16}})", text):
            seen.setdefault(m.group(1), {}).setdefault(m.group(2), set()).add(idx)
    for num, labels in seen.items():
        if len(labels) > 1:
            detail = " / ".join(f"「{k}」p{','.join(map(str, sorted(v)))}" for k, v in list(labels.items())[:3])
            warn(f"番号つきの項目名が食い違う: {num} が {detail}（§7.13・§7.33: 同じ番号は同じ項目名で。別の意味なら番号を分ける）")
            return


def page_number(fragment):
    """フッターに印字したページ番号（基本パーツ集 .foot の span／追加パーツ集 .footer の最後の span）"""
    m = re.search(r'<div class="foot">.*?<span>\s*(\d+)\s*</span>', fragment, re.S) or \
        re.search(r'<footer class="footer">(?:(?!</footer>).)*<span>\s*(\d+)\s*</span>\s*</footer>', fragment, re.S)
    return int(m.group(1)) if m else None


def _part_type_ids():
    """追加パーツ集の英語の型ID（右上ラベル「05  chart_insight」の chart_insight）。テンプレが無い環境では空。"""
    fp = Path(__file__).resolve().parent.parent / "templates" / "freeform_parts_more_16x9.html"
    if not fp.exists():
        return set()
    return set(re.findall(r'class="kicker">\s*\d{2}\s+([a-z][a-z0-9_]*)\s*<', fp.read_text(encoding="utf8", errors="ignore")))


_TYPE_IDS = _part_type_ids()
# パーツ集にある型IDだけを見る（表の「05」「user_id」のような列の並びを型IDと取り違えない）
TYPE_ID = re.compile(r"(?<![\w/.])\d{2}\s+(?:" + "|".join(sorted(map(re.escape, _TYPE_IDS), key=len, reverse=True)) + r")(?![\w.])") if _TYPE_IDS else None


def check_type_id(idx, text):
    """パーツの英語の型ID（「05 chart_insight」など）が右上のラベルや本文に残っていないか"""
    m = TYPE_ID.search(text) if TYPE_ID else None
    if m:
        fail(f"p{idx}: パーツの型ID「{m.group(0)}」が残っている（右上のラベルはそのページの話題を日本語で書く。§4.13）")


def check_page_numbers(nums):
    """§4.45: 印字したページ番号は連番（表紙・章扉・裏表紙は番号を振らないので除く）"""
    got = [(i, n) for i, n in nums if n is not None]
    for (i0, a), (i1, b) in zip(got, got[1:]):
        if b != a + 1:
            fail(f"p{i1}: ページ番号が連番でない（前のページ {a} → {b}）（ページを足す・消すしたあとに振り直す。§4.45）")
            return


def header_label(fragment):
    """本文ページの右上ラベルの有無。基本パーツは `.bar`/`.hdr` の中の `.date`、
    追加パーツは `.kicker` がラベル。どちらも同じ「ラベル」として扱い、文字が入っているかで有無を見る。
    タイトルより上だけを見る（グラフの棒 `.bar` などを拾わない）"""
    head = re.split(r'<h1\b|class="[^"]*\b(?:ttl|title|msg)\b', fragment, maxsplit=1)[0]
    m = re.search(r'<(\w+)[^>]*class="(?:[^"]*\s)?(?:kicker|date)(?:\s[^"]*)?"[^>]*>(.*?)</\1>', head, re.S)
    text = re.sub(r"<[^>]+>", "", m.group(2)).strip() if m else ""
    return "あり" if text else "なし"


def check_header_consistency(heads):
    """§4.19・§4.62: 本文ページの右上ラベル（タグチップ）の有無と作りをデッキで揃える"""
    body = [(i, h) for i, h in heads if h is not None]
    if len(body) < 4:
        return
    from collections import Counter
    major, cnt = Counter(h for _, h in body).most_common(1)[0]
    odd = [i for i, h in body if h != major]
    if odd and cnt >= len(body) * 0.6:
        warn(f"p{','.join(map(str, odd[:6]))}: 右上のラベルの有無・作りがほかの本文ページと違う（§4.19: 見出しまわりの作りはデッキで揃える）")


def check_exec_summary(idx, title, fragment):
    """§7.16 エグゼクティブサマリーはページタイトルの羅列にせず、主文＋インデントした詳細の階層ブレットで書く"""
    if not re.search(r"エグゼクティブサマリー|エグゼクティブ・サマリー|Executive Summary", title, re.I):
        return
    if not re.search(r"<li\b(?:(?!</li>).)*?<(?:ul|ol)\b", fragment, re.S):
        warn(f"p{idx}: エグゼクティブサマリーに入れ子のブレットが無い（§7.16: 各行を主文＋インデントした詳細2〜3本で書く。ページタイトルの羅列にしない）")


def check_html(path):
    global STRICT_LEN
    STRICT_LEN = False
    html = Path(path).read_text(encoding="utf8", errors="ignore")
    if not ("338.67mm" in html and "190.5mm" in html):
        fail("16:9 サイズ（338.67mm×190.5mm）が CSS に見当たらない（A4 / 297×167 は旧仕様）")
    if re.search(r"@page\s*{[^}]*297mm", html):
        fail("@page が A4 横のまま")
    for c in OLD_COLORS:
        if has_color(c, html):
            fail(f"禁止色 #{c}")
    # 角丸
    for m in re.finditer(r"([^{}]{0,80}){[^}]*?border-radius\s*:\s*(\d+(?:\.\d+)?)(px|mm|rem|em)", html):
        sel, v, u = m.group(1), float(m.group(2)), m.group(3)
        if re.search(r"pill|chip|tag|badge|dot", sel, re.I):
            continue  # 小ピルは丸可
        px = v * {"px": 1, "mm": 3.78, "rem": 16, "em": 16}[u]
        if px >= 4:
            fail(f"border-radius {v}{u} on `{sel.strip()[-40:]}`（角丸禁止。小ピル以外は直角）")
            break
    # サブタイトル・図表ラベル
    for cls in ["figttl", "subtitle"]:
        if re.search(r'class="[^"]*\b' + cls + r'\b', html):
            fail(f"サブタイトル/図表ラベル系クラス `.{cls}` が残っている")
    # 表紙・章扉・裏表紙（class に cover/chap/back）のサブタイトルは数えない
    frags, fcls = _slides(html)
    inner = "".join(f for f, c in zip(frags, fcls) if not {"cover", "chap", "back"} & set(c.split())) if frags else html
    for cls in ["sub", "lead", "caption"]:
        n = len(re.findall(r'class="[^"]*\b' + cls + r'\b', inner))
        if n:
            warn(f"`.{cls}` ×{n} — 表紙サブタイトルなら可。コンテンツスライドのタイトル直下なら禁止（目視確認）")
    if re.search(r'<span class="ac">', html):
        fail("タイトル内の色分け <span class=\"ac\"> が残っている")
    # 表ヘッダー
    th = re.search(r"\bth\s*{[^}]*font-size\s*:\s*(\d+)px", html)
    td = re.search(r"\btd\s*{[^}]*font-size\s*:\s*(\d+)px", html)
    if th and td and int(th.group(1)) < int(td.group(1)) + 2:
        fail(f"表ヘッダー {th.group(1)}px が本文 {td.group(1)}px +2pt 未満")
    elif th and td and int(th.group(1)) < int(td.group(1)) + 3:
        warn(f"表ヘッダー {th.group(1)}px は本文 {td.group(1)}px +3px 未満（§6: 見出しは本文より+3〜4pt 大きくするのが目安）")
    if re.search(r"\bth\s*{[^}]*color\s*:\s*#?(9[0-9a-f]{5}|a[0-9a-f]{5}|b[0-9a-f]{5}|c[0-9a-f]{5}|888|999|aaa|bbb|ccc|gr[ae]y)\b", html, re.I):
        warn("表ヘッダーが薄グレー（§6: 見出しは本文と同じ濃色）")
    check_kicker_and_conclusion(html)
    # ハーベイボール ¾ の描画（中心点の無い多角形は斜めに欠けた形になる）
    if re.search(r"\.q3\s+i\s*{[^}]*clip-path\s*:\s*polygon\(\s*50%\s+0\s*,", html):
        fail("ハーベイボール ¾ の clip-path に中心点（50% 50%）が無い。左上が斜めに欠けた形になる（polygon(50% 50%,50% 0,100% 0,100% 100%,0 100%,0 50%) に直す）")
    if not re.search(r"<meta[^>]+charset\s*=\s*[\"']?utf-?8", html, re.I):
        fail('<meta charset="utf-8"> が無い（Windows のブラウザで文字化けする — slide-rules §8）')
    if re.search(r"\bth\s*{[^}]*font-weight\s*:\s*(400|normal|300)", html):
        fail("表ヘッダーが細字")
    if re.search(r"tr:nth-child\((even|odd)\)", html):
        fail("ゼブラ縞が残っている")
    # 枠線ルール（slide-rules §5.3-5.4 / §6）
    if re.search(r"\btd\b[^{}]*{[^}]*border-bottom\s*:", html) and not re.search(
            r"tr:last-child[^{}]*{[^}]*border(-bottom)?\s*:\s*(0|none)", html):
        fail("最終行の罫線が消えていない（`tr:last-child td{border-bottom:0}` を追加 — 行き先のない罫線禁止）")
    body_html = html.split("</style>", 1)[1] if "</style>" in html else html
    used_classes = set(re.findall(r'class="([^"]*)"', body_html))
    used_tokens = set(tok for cl in used_classes for tok in cl.split())
    for m in re.finditer(r"([^{}]{0,80}){([^}]*)}", html):
        sel, body = m.group(1), m.group(2)
        if re.search(r"pill|chip|tag|badge|dot|legend", sel, re.I):
            continue
        # 本文で使っていないクラスの規則は対象外（パーツ集のCSSをまとめて取り込んだデッキで誤検知しない）
        cls_in_sel = re.findall(r"\.([A-Za-z0-9_-]+)", sel)
        if cls_in_sel and cls_in_sel[-1] not in used_tokens:  # 主語（末尾のクラス）が本文に無ければ対象外
            continue
        has_fill = re.search(r"background(-color)?\s*:\s*(?!none|transparent)#?\w", body)
        has_border = re.search(r"border\s*:\s*(?!0|none)\d", body)
        if has_fill and has_border and re.search(r"card|box|pillar|mem|step|stat", sel, re.I):
            warn(f"塗りありボックスに枠線: `{sel.strip()[-40:]}`（塗りカードは border:0 — slide-rules §5.3）")
    # タイトル抽出
    titles = []
    shapes = []
    slides, slide_classes = _slides(html)
    nums, heads = [], []
    for i, s in enumerate(slides, 1):
        vis = _visible(s)
        check_garbled(i, vis)
        check_work_files(i, vis)
        if not TEMPLATE_MODE:
            check_type_id(i, vis)
        m = re.search(r"<h1[^>]*>(.*?)</h1>", s, re.S) or re.search(r'class="[^"]*\b(?:ttl|title|msg)\b[^"]*"[^>]*>(.*?)</', s, re.S)
        raw_t = m.group(1) if m else ""
        t = re.sub(r"<[^>]+>", "", raw_t).strip()
        t = re.sub(r"\s+", " ", t)
        titles.append(t)
        cls = (slide_classes[i - 1] if i - 1 < len(slide_classes) else "").split()
        nums.append((i, page_number(s)))
        others = [x for x in _leaf_texts(s) if x != t and not re.fullmatch(r"\d{1,3}", x)]
        APPX_PAGES[i] = (appendix_tag_html(s, m, t), sum(_fwlen(x) for x in others),
                         bool(re.search(r"<(?:table|svg|canvas|video|iframe)\b", s, re.I))
                         or len(re.findall(r"<img\b", s, re.I)) >= 2)  # 画像1枚はロゴとみなす
        # 最終ページは、cover/back の印が無くても、タイトルが無いか締めのあいさつだけのときに限って裏表紙とみなす
        # （裏表紙の無いデッキの最終ページの比較表の「当社」を見本の文言と取り違えない）
        is_back = "back" in cls or (i == len(slides) and (not t or bool(CLOSING_TITLE.search(t))))
        if not TEMPLATE_MODE and ("cover" in cls or is_back):
            left = [x for x in _leaf_texts(s) if re.fullmatch(r"当社|会社名|連絡先|お問い合わせ\s*当社|資料名|担当者名|表紙|裏表紙|章扉", x)]
            if left:
                fail(f"p{i}: 表紙・裏表紙に見本の文言が残っている「{left[0]}」（作成者・連絡先は実物にし、「表紙」のような型の名前のラベルは消す）")
        edge = bool({"cover", "chap", "back"} & set(cls))
        if "std" in cls:
            # 正本から差し込んだ定型ページ（§4.99）は文章・版面の検査を飛ばす
            shapes.append("")
            continue
        if not edge:
            heads.append((i, header_label(s)))
        # <br> で意味の切れ目に明示改行済み（§2.13）なら、各行が1行目安に収まる限り2行想定の WARN は出さない
        lines = [re.sub(r"\s+", " ", re.sub(r"<[^>]+>", "", x)).strip() for x in re.split(r"<br\s*/?>", raw_t)]
        explicit_break = len(lines) > 1 and all(_fwlen(x) <= TITLE_MAX for x in lines)
        closing = edge or (i == len(slides) and bool(CLOSING_TITLE.search(t)))
        check_title(i, t, explicit_break, closing, edge)
        check_count_match(i, t, re.sub(r"<[^>]+>", " ", s))
        leaves = _leaf_texts(s)
        check_body_placeholders(i, leaves, t)
        check_framework_and_commas(i, leaves, vis)
        check_dash_link(i, leaves)
        check_exec_summary(i, t, s)
        if "cover" not in cls:
            check_multi_sentence(i, leaves)
            check_trailing_period(i, [x for x in leaves if x != t])
        check_axis_names_html(i, s)
        if not closing:
            check_body_desu(i, leaves, t)
        check_single_item_lists(i, s)
        check_typed_bullets(i, s)
        shapes.append(page_shape(s))
    check_page_refs([(i, re.sub(r"<[^>]+>", " ", s)) for i, s in enumerate(slides, 1)])
    PAGES[:] = [(i, re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", s, flags=re.S))))
                for i, s in enumerate(slides, 1)]
    if not slides:
        warn("`.slide` 要素が見つからない（タイトル検査スキップ）")
    check_shape_runs(shapes)
    if slides and not TEMPLATE_MODE:
        check_page_numbers(nums)
        check_header_consistency(heads)
        check_numbered_labels([(i, re.sub(r"<[^>]+>", " ", s)) for i, s in enumerate(slides, 1)])
    if re.search(r"word-break\s*:\s*keep-all", html):
        warn("word-break:keep-all（括弧・句読点でしか折れず行が早く切れる。日本語は word-break:auto-phrase と lang=\"ja\" — §2.13）")
    check_static_ui(html)
    check_mixed_parts(html, slide_classes)
    check_production_meta(slides if slides else [html])
    if slides:
        check_terms([(i, re.sub(r"<[^>]+>", " ", s)) for i, s in enumerate(slides, 1)])
        check_number_consistency([(i, re.sub(r"<[^>]+>", " ", s)) for i, s in enumerate(slides, 1)])
        check_ai_smell([(i, re.sub(r"<[^>]+>", "\n", s)) for i, s in enumerate(slides, 1)])
        check_forbidden([(i, re.sub(r"<[^>]+>", " ", s)) for i, s in enumerate(slides, 1)], html)
    else:
        check_terms([(1, re.sub(r"<[^>]+>", " ", html))])
        check_ai_smell([(1, re.sub(r"<[^>]+>", "\n", html))])
    return titles


NUM_FACT = re.compile(r"(?<![\d.,])(\d{1,3}(?:[,，]\d{3})+|\d+(?:\.\d+)?)\s*(%|％|ポイント|倍|億円|万円|千円|円|億|万|件|人|社|時間|ドル)")


def _norm_num(x):
    return x.replace(",", "").replace("，", "")


SRC_NUM = re.compile(r"(?<![\d.])\d+(?:\.\d+)?(?![\d])")


def _src_numbers(text):
    """sources.md の数字を、前後が数字でない境界で切り出した集合にする（部分一致で「2025」が「5」を通さないように）。
    桁区切りのカンマは外し、小数の末尾の 0 は落とした形も入れる（12.50 ↔ 12.5）"""
    text = re.sub(r"(?<=\d)[,，](?=\d{3}(?!\d))", "", text)
    out = set()
    for t in SRC_NUM.findall(text):
        out.add(t)
        if "." in t:
            out.add(t.rstrip("0").rstrip("."))
    return out


SLIDE_OPEN = re.compile(r'<(?:section|div)[^>]*class="((?:[^"]*\bslide\b[^"]*|s|s [^"]*))"[^>]*>')


def _slides(html):
    """スライドの区切り。(各スライドの断片のリスト, 各スライドの class のリスト) を返す。
    スライドを数える検査はすべてこれを使い、ページ番号がずれないようにする"""
    parts = SLIDE_OPEN.split(html)
    # キャプチャした class も返るので [前置き, class1, 断片1, class2, 断片2, ...] になる
    return parts[2::2], parts[1::2]


def _visible(frag):
    """script・style・コメントを除いた見える文字（タグは空白に置き換える）"""
    return re.sub(r"<[^>]+>", " ", re.sub(r"<(script|style)[^>]*>.*?</\1>|<!--.*?-->", " ", frag, flags=re.S))


def check_sources(html_path, sources_path):
    """数字の出典の控え（sources.md）と照合する。デッキに出てくる数字（単位つき）が控えに無ければ WARN（SKILL.md 手順2）"""
    sp = Path(sources_path)
    if not sp.exists():
        fail(f"出典の控え {sources_path} が無い（SKILL.md 手順2: 数字は調べて控えてから使う）")
        return
    src = _src_numbers(sp.read_text(encoding="utf8", errors="ignore"))
    html = Path(html_path).read_text(encoding="utf8", errors="ignore")
    missing = []
    frags, classes = _slides(html)
    for i, frag in enumerate(frags, 1):
        if "std" in (classes[i - 1] if i - 1 < len(classes) else "").split():
            continue  # 正本から差し込んだ定型ページ（§4.99・会社概要など）は照合しない。ほかの文章検査と揃える
        vis = _visible(frag)
        for m in NUM_FACT.finditer(vis):
            n = _norm_num(m.group(1))
            if n not in src and (n.rstrip("0").rstrip(".") if "." in n else n) not in src:
                missing.append(f"p{i}「{m.group(0).strip()}」")
    if missing:
        uniq = list(dict.fromkeys(missing))
        warn(f"出典の控えに無い数字 ×{len(uniq)}: {' / '.join(uniq[:8])}（{sources_path} に出典を控えるか、試算なら前提と算式を書く。見つからない数字は「要確認」に — §2.23）")


def usage_error(msg):
    """引数・設定ファイルの誤りは traceback を出さずに理由を出して終える"""
    print(f"check_deck.py: {msg}", file=sys.stderr)
    sys.exit(2)


def load_colors(path):
    """1行1色。行頭の6桁の HEX（# は付けても付けなくてもよい）を読む。後ろの説明（「#C5A681 旧グレー」）は無視する。
    # で始まり色で始まらない行はコメント、空行は読み飛ばす。色として読めない行は WARN にして読み飛ばす"""
    p = Path(path).expanduser()
    if not p.is_file():
        usage_error(f"--old-colors のファイルが見つからない: {path}")
    out = []
    for n, line in enumerate(p.read_text(encoding="utf-8-sig").splitlines(), 1):
        s = line.strip()
        if not s:
            continue
        m = re.match(r"#?([0-9A-Fa-f]{6})(?:[0-9A-Fa-f]{2})?(?![0-9A-Za-z])", s)  # 8桁（透明度つき）は先頭6桁
        if m:
            out.append(m.group(1).upper())
        elif re.match(r"#?[0-9A-Fa-f]{3,4}(?![0-9A-Za-z])", s):
            warn(f"{p.name} {n}行目: 3桁の HEX は色として読めないので読み飛ばした「{s[:30]}」（#8A8 → #88AA88 のように6桁で書く）")
        elif not s.startswith("#"):
            warn(f"{p.name} {n}行目: 色として読めないので読み飛ばした「{s[:30]}」（1行1色・HEX 6桁で書く）")
    return out


def has_color(c, text):
    """6桁の HEX がそのまま現れるか（大文字小文字は問わない。8桁の透明度つき「#8a8f94cc」も同じ色とみなす。
    それ以外で前後が16進の文字なら別の値の一部とみなす）"""
    return re.search(r"(?<![0-9A-Fa-f])" + re.escape(c) + r"(?:[0-9A-Fa-f]{2})?(?![0-9A-Fa-f])", text, re.I) is not None


def _cell_text(cell_html):
    t = re.sub(r"<[^>]+>", "", cell_html)
    return re.sub(r"(&nbsp;|&#160;|\s|　)+", "", t)


def check_axis_names_html(idx, slide_html):
    """§4.102: 見出し行（th だけの行）の左端セルが空欄なら FAIL。他の見出しセルに文字があるときだけ判定"""
    n = 0
    for table in re.findall(r"<table\b.*?</table>", slide_html, re.S | re.I):
        for tr in re.findall(r"<tr\b[^>]*>(.*?)</tr>", table, re.S | re.I):
            cells = re.findall(r"<(t[hd])\b[^>]*>(.*?)</\1>", tr, re.S | re.I)
            if len(cells) < 2 or any(tag.lower() != "th" for tag, _ in cells):
                continue
            if not _cell_text(cells[0][1]) and any(_cell_text(c) for _, c in cells[1:]):
                n += 1
    if n:
        fail(f"p{idx}: 表の見出し行の左端セルが空欄 ×{n}（§4.102: フェーズ帯・ステップ帯にも軸の名前を書く）")


def _shape_type(sh):
    """shape_type（種類の決まらない図形・プレースホルダーでは例外になるので None）"""
    try:
        return sh.shape_type
    except Exception:
        return None


def _pptx_head_row(tbl, r, row):
    """PPTX の表で、この行が見出しの行か。表の書式で先頭行を見出しにしている（firstRow）か、
    文字のあるセル（左端以外）がすべて太字のときだけ見出しとみなす。見出しの無い表の1行目は対象外"""
    if r == 0 and tbl._tbl.tblPr is not None and tbl._tbl.tblPr.get("firstRow") in ("1", "true"):
        return True
    runs = [[x for p in c.text_frame.paragraphs for x in p.runs if x.text.strip()]
            for c in list(row.cells)[1:] if c.text.strip()]
    return bool(runs) and all(rs and all(x.font.bold for x in rs) for rs in runs)


def check_axis_names_pptx(idx, slide):
    """§4.102 の PPTX 版: 表の先頭の見出し帯（見出しの行が続くまで・最大3段）で左端セルが空欄なら FAIL"""
    for sh in slide.shapes:
        if not getattr(sh, "has_table", False):
            continue
        tbl = sh.table
        if len(tbl.rows) < 2 or len(tbl.columns) < 2:
            continue
        for r, row in enumerate(list(tbl.rows)[:3]):
            cells = list(row.cells)
            if cells[0].is_spanned or cells[0].text.strip() or not _pptx_head_row(tbl, r, row):
                break
            if any(c.text.strip() for c in cells[1:]):
                fail(f"p{idx}: 表の見出し行の左端セルが空欄（§4.102: フェーズ帯・ステップ帯にも軸の名前を書く）")
                return


PAGES = []  # [(idx, 本文テキスト)] check_pptx / check_html が詰める
APPX_PAGES = {}  # {idx: (付録のタグチップがあるか, タイトル以外の文字数, 表・図・画像があるか)} §4.100 用
APPX_LABEL = r"(?:appendix|付録|参考資料|補足資料)"
APPX_SEP = r"[\s:：／/|｜・\-－—（(【\[「]"
# 付録ページのタイトル: 「Appendix」「付録」の直後が終わり・区切り記号・空白のとき（「付録の手順書を改め…」は本編）
APPX_TITLE_RE = re.compile(r"^\s*(?:appendix|付録)(?=$|" + APPX_SEP + ")", re.I)
# 付録のタグチップ: 要素の文字が「付録」「Appendix」（「Appendix A-1」「付録 2」まで）だけのもの
APPX_TAG_RE = re.compile(r"^\s*(?:appendix|付録)(?:\s*[A-Za-z0-9][A-Za-z0-9.\-]*)?\s*$", re.I)
# アジェンダ・目次のページ（項目の「Appendix」は付録のタグではない）
AGENDA_TITLE_RE = re.compile(r"(?i)agenda|contents|アジェンダ|目次|本日の(?:流れ|内容)|資料の構成")
# タグチップらしい要素のクラス（「es-tag」のような部品内のラベルは含めない）
TAG_CLASS_RE = re.compile(r'<(\w+)\b[^>]*\bclass="(?:[^"]*\s)?(?:tag|chip|badge|kicker|eyebrow)(?:\s[^"]*)?"[^>]*>(.*?)</\1>', re.S)


def appendix_tag_html(slide, title_match, title):
    """付録のタグチップがあるか。タグとみなすのは、クラスが tag・chip 等の要素と、タイトル（h1）より前の要素だけ。
    アジェンダの「<li>Appendix</li>」や表のセルの「付録」は数えない"""
    if AGENDA_TITLE_RE.search(title or ""):
        return False
    cands = [re.sub(r"\s+", " ", re.sub(r"<[^>]+>", "", m.group(2))).strip() for m in TAG_CLASS_RE.finditer(slide)]
    if title_match:
        cands += _leaf_texts(slide[:title_match.start()])
    return any(APPX_TAG_RE.match(x) for x in cands)


def appendix_tag_pptx(slide, title_shape, title, sw, sh_):
    """付録のタグチップがあるか（PPTX）。上端付近（高さの上18%）にある小さい図形の文字だけを見る。
    本文の「Appendix」だけの図形（アジェンダの項目など）は数えない"""
    if AGENDA_TITLE_RE.search(title or "") or not sw or not sh_:
        return False
    for sh in slide.shapes:
        if sh is title_shape or not sh.has_text_frame or sh.top is None or sh.height is None or sh.width is None:
            continue
        if (sh.top < sh_ * 0.18 and sh.height <= sh_ * 0.1 and sh.width <= sw * 0.3
                and APPX_TAG_RE.match(sh.text_frame.text.strip())):
            return True
    return False


# 文になっている（区切りのラベルではなく、付録ページのメッセージ）
SENTENCE_RE = re.compile(r"[、。]|[がを]|(?:た|だ|る|ます|です|ない|れる)$")
DIVIDER_MAX_BODY = 100  # 区切りのスライドのタイトル以外の文字数の上限（中身1行＋資料名・ページ番号）


def divider_label(t):
    """区切りのタイトルらしいか: 「Appendix」「付録 A」「Appendix｜部門別の数字」のように、ラベルと短い名詞句だけ"""
    m = re.match(r"^\s*" + APPX_LABEL + r"(.*)$", t or "", re.I | re.S)
    if not m:
        return False
    rest = m.group(1)
    if rest.strip() and not re.match(APPX_SEP + r"|[A-Za-z0-9]", rest):
        return False  # 「付録の手順書…」のように語が続く
    rest = re.sub(r"^" + APPX_SEP + r"+", "", rest.strip())
    return _fwlen(rest) <= 24 and not SENTENCE_RE.search(rest)


def check_appendix(titles):
    """§4.100: 付録らしきページがあるのに、その手前に区切りスライドがなければ FAIL。
    付録ページ = タイトルが「Appendix」「付録」で始まる（直後が区切り記号・空白）か、付録のタグチップがあるページ。
    区切り = タイトルが Appendix 等のラベル（短い名詞句まで）で、本文が短く、表・図が無いページ"""
    if not PAGES:
        return
    ttl = lambda i: (titles[i - 1] or "").strip() if i - 1 < len(titles) else ""
    info = lambda i: APPX_PAGES.get(i, (False, 0, False))
    first = next((i for i, _ in PAGES if i > 1 and (info(i)[0] or APPX_TITLE_RE.search(ttl(i)))), None)
    if first is None:
        return
    body = dict(PAGES)

    def is_div(i):
        _, n, rich = info(i)
        if rich or n > DIVIDER_MAX_BODY:
            return False
        if ttl(i):
            return divider_label(ttl(i))
        # タイトルを持たない扉（h1 以外で組んだもの）は本文の先頭で見る
        return bool(re.match(r"^\s*" + APPX_LABEL, body.get(i, ""), re.I))

    # first 自身も、区切りの形（短く・表や図が無い）のときだけ区切りに数える
    if not any(is_div(i) for i in range(2, first + 1)):
        fail(f"p{first}: 付録（Appendix）の手前に区切りスライドがない（§4.100: 本編と付録の間に「Appendix」＋中身1行の扉を1枚挟む）")


def main():
    global TEMPLATE_MODE
    args = sys.argv[1:]
    global FORBIDDEN_TERMS, PART_TYPE_NAMES, XML_ONLY, HOUSE
    brand = None
    if "--brand" in args:
        k = args.index("--brand")
        if k + 1 >= len(args) or args[k + 1].startswith("--"):
            usage_error("--brand の値がない（いま受け付ける値は none だけ: --brand none）")
        brand = args[k + 1]
        if brand not in BRANDS:
            usage_error(f"--brand の値「{brand}」は使えない（いま受け付ける値は {' / '.join(BRANDS)}）")
        del args[k:k + 2]
    colors_path = None
    if "--old-colors" in args:
        k = args.index("--old-colors")
        if k + 1 >= len(args) or args[k + 1].startswith("--"):
            usage_error("--old-colors のファイルの指定がない（--old-colors 色の一覧.txt）")
        colors_path = args[k + 1]
        del args[k:k + 2]
    if brand == "none" and colors_path:
        print(f"--brand none のため --old-colors {colors_path} は使わない（レガシー色の検査を外している）", file=sys.stderr)
    if brand != "none":
        if colors_path:
            OLD_COLORS[:] = load_colors(colors_path)
        elif (LOCAL_DIR / "old-colors.txt").exists():
            OLD_COLORS[:] = load_colors(LOCAL_DIR / "old-colors.txt")
    if "--xml-only" in args:
        XML_ONLY = True
        args.remove("--xml-only")
    if "--house" in args:
        k = args.index("--house")
        import json
        HOUSE = json.loads(Path(args[k + 1]).read_text(encoding="utf8"))
        del args[k:k + 2]
    if "--template" in args:
        TEMPLATE_MODE = True
        args.remove("--template")
    if "--forbid" in args:
        k = args.index("--forbid")
        FORBIDDEN_TERMS = load_forbidden(args[k + 1])
        del args[k:k + 2]
    sources = None
    if "--sources" in args:
        k = args.index("--sources")
        sources = args[k + 1]
        del args[k:k + 2]
    PART_TYPE_NAMES = _part_type_names()
    if not args:
        sys.exit(__doc__)
    p = args[0]
    titles = check_pptx(p) if p.lower().endswith((".pptx", ".potx")) else check_html(p)
    if not TEMPLATE_MODE:
        check_title_variety(titles)
        check_count_titles(titles)
        check_appendix(titles)
    if sources and not p.lower().endswith((".pptx", ".potx")):
        check_sources(p, sources)
    print("=== タイトル一覧（上から通し読みしてストーリーが繋がるか確認） ===")
    for i, t in enumerate(titles, 1):
        print(f"{i:>3}  {t or '(なし)'}")
    print()
    for w in warns:
        print("WARN ", w)
    for f in fails:
        print("FAIL ", f)
    print(f"\n{len(fails)} FAIL / {len(warns)} WARN")
    sys.exit(1 if fails else 0)


if __name__ == "__main__":
    main()
