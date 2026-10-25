"""機械チェックの自己テスト（標準ライブラリのみ）。

指摘を機械チェックに足したら、ここに「直していない版で FAIL が出る」「直した版で出ない」の
両方を1件ずつ足す。直していない版で発火しないチェックは、測れていないのと同じ。

  python3 -m unittest discover -s tests
"""
import re
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
CHECK = ROOT / "scripts" / "check_deck.py"
GOOD = (ROOT / "tests" / "fixtures" / "good_deck.html").read_text(encoding="utf8")
# 利用者の手元の local/old-colors.txt に左右されないよう、レガシー色の検査は外して走らせる
NO_BRAND = ("--brand", "none")


def run(html, *extra):
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "deck.html"
        p.write_text(html, encoding="utf8")
        extra = [str(Path(d) / a[1:]) if a.startswith("@") else a for a in extra]
        for a in extra:
            if a.endswith(".txt"):
                Path(a).write_text("# 架空の禁止語\nサンプル商事\nre:PJ-\\d{3}\n", encoding="utf8")
        r = subprocess.run([sys.executable, str(CHECK), str(p), *NO_BRAND, *extra], capture_output=True, text=True)
    return r.returncode, [l for l in r.stdout.splitlines() if l.startswith("FAIL")]


def warns(html):
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "deck.html"
        p.write_text(html, encoding="utf8")
        r = subprocess.run([sys.executable, str(CHECK), str(p), *NO_BRAND], capture_output=True, text=True)
    return [l for l in r.stdout.splitlines() if l.startswith("WARN")]


def inject(old, new):
    assert old in GOOD, old
    return GOOD.replace(old, new, 1)


def run_files(html, files, *extra):
    """files: {名前: 中身} を一時フォルダに置いてから検査する。extra の @名前 はその置き場所に読み替える"""
    with tempfile.TemporaryDirectory() as d:
        for name, body in files.items():
            (Path(d) / name).write_text(body, encoding="utf8")
        p = Path(d) / "deck.html"
        p.write_text(html, encoding="utf8")
        extra = [str(Path(d) / a[1:]) if a.startswith("@") else a for a in extra]
        r = subprocess.run([sys.executable, str(CHECK), str(p), *extra], capture_output=True, text=True)
    return r.returncode, [l for l in r.stdout.splitlines() if l.startswith("FAIL")]


class GoodDeck(unittest.TestCase):
    def test_good_deck_passes(self):
        code, fails = run(GOOD)
        self.assertEqual((code, fails), (0, []))

    def test_templates_pass_in_template_mode(self):
        for f in ("freeform_parts_16x9.html", "freeform_parts_more_16x9.html"):
            r = subprocess.run([sys.executable, str(CHECK), str(ROOT / "templates" / f), "--template", *NO_BRAND],
                               capture_output=True, text=True)
            self.assertEqual(r.returncode, 0, f + r.stdout[-400:])


class Placeholders(unittest.TestCase):
    def test_body_placeholder_fires(self):
        code, fails = run(inject("承認待ちは平均3日", "Text 1"))
        self.assertEqual(code, 1)
        self.assertTrue(any("本文にテンプレのプレースホルダー" in f for f in fails), fails)

    def test_part_label_chip_fires(self):
        code, fails = run(inject('<div class="date">現状</div>', '<div class="date">パーツ06｜前提→帰結の2カラム</div>'))
        self.assertTrue(any("プレースホルダー" in f for f in fails), fails)

    def test_type_name_title_fires(self):
        code, fails = run(inject("<h1>1部門で2か月試行し、効果を確かめてから全社へ広げる</h1>", "<h1>軸のある表</h1>"))
        self.assertTrue(any("型名のまま" in f for f in fails), fails)


class OneSentencePerBlock(unittest.TestCase):
    def test_two_sentences_in_card_fires(self):
        code, fails = run(inject("• 例外だけ承認者に回す", "• 例外だけ承認者に回す。規程外の申請は差し戻す"))
        self.assertTrue(any("2文以上" in f for f in fails), fails)

    def test_period_inside_brackets_is_ignored(self):
        code, fails = run(inject("• 例外だけ承認者に回す", "• 例外（金額超過。科目不明）だけ承認者に回す"))
        self.assertFalse(any("2文以上" in f for f in fails), fails)

    def test_single_sentence_with_trailing_period_passes(self):
        code, fails = run(inject("• 例外だけ承認者に回す", "• 例外だけ承認者に回す。"))
        self.assertFalse(any("2文以上" in f for f in fails), fails)


def run_warns(html):
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "deck.html"
        p.write_text(html, encoding="utf8")
        r = subprocess.run([sys.executable, str(CHECK), str(p), *NO_BRAND], capture_output=True, text=True)
    return [l for l in r.stdout.splitlines() if l.startswith("WARN")]


class NumberConsistency(unittest.TestCase):
    """slide-rules §7.6 数値の平仄: 同じ指標がページ間で違う値なら WARN。"""

    def deck(self, p2, p5):
        return inject("承認待ちは平均3日", p2).replace("営業部", p5, 1)

    def test_same_label_different_value_fires(self):
        warns = run_warns(self.deck("承認者数は12人", "承認者数は15人"))
        self.assertTrue(any("数値の平仄疑い" in w and "承認者数" in w for w in warns), warns)

    def test_same_value_in_other_notation_passes(self):
        warns = run_warns(self.deck("承認者数は1.2万人", "承認者数 12,000人"))   # 桁の書き方が違っても値が同じなら可
        self.assertFalse(any("数値の平仄疑い" in w for w in warns), warns)

    def test_amount_with_scale_fires(self):
        warns = run_warns(self.deck("売上高は120億円", "売上高 118億円"))   # 金額・割合など単位を問わず比べる
        self.assertTrue(any("数値の平仄疑い" in w and "売上高" in w for w in warns), warns)

    def test_same_amount_in_other_scale_passes(self):
        warns = run_warns(self.deck("売上高は1.2億円", "売上高 120,000,000円"))
        self.assertFalse(any("数値の平仄疑い" in w for w in warns), warns)

    def test_float_rounding_does_not_fire(self):
        warns = run_warns(self.deck("売上高は1.1億円", "売上高 110,000,000円"))   # 1.1×1e8 の float 誤差を食い違いにしない
        self.assertFalse(any("数値の平仄疑い" in w for w in warns), warns)

    def test_different_year_passes(self):
        warns = run_warns(self.deck("2024年の承認者数は12人", "2026年の承認者数は15人"))   # 時点が違えば別の指標
        self.assertFalse(any("数値の平仄疑い" in w for w in warns), warns)


class ForbiddenTerms(unittest.TestCase):
    def test_term_in_body_fires_without_echoing_it(self):
        code, fails = run(inject("営業部", "サンプル商事の営業部"), "--forbid", "@terms.txt")
        self.assertTrue(any("禁止語" in f for f in fails), fails)
        self.assertFalse(any("サンプル商事" in f for f in fails), fails)

    def test_term_in_html_comment_fires(self):
        code, fails = run(inject("</main>", "<!-- PJ-123 向けの下書き --></main>"), "--forbid", "@terms.txt")
        self.assertTrue(any("コメント/属性" in f for f in fails), fails)

    def test_no_list_no_check(self):
        code, fails = run(inject("営業部", "サンプル商事の営業部"))
        self.assertEqual(code, 0)


def run_out(html):
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "deck.html"
        p.write_text(html, encoding="utf8")
        r = subprocess.run([sys.executable, str(CHECK), str(p), *NO_BRAND], capture_output=True, text=True)
    return r.stdout


class Charset(unittest.TestCase):
    def test_missing_charset_fires(self):
        html = re.sub(r"<meta[^>]*charset[^>]*>", "", GOOD, count=1)
        self.assertNotEqual(html, GOOD)
        code, fails = run(html)
        self.assertTrue(any("charset" in f for f in fails), fails)

    def test_good_deck_has_charset(self):
        code, fails = run(GOOD)
        self.assertFalse(any("charset" in f for f in fails), fails)


class TitleConnector(unittest.TestCase):
    T = "<h1>1部門で2か月試行し、効果を確かめてから全社へ広げる</h1>"

    def test_connector_start_warns(self):
        out = run_out(inject(self.T, "<h1>まずは1部門で2か月試行し、効果を確かめてから全社へ広げる</h1>"))
        self.assertIn("接続詞で始まる", out)

    def test_tsugini_as_verb_phrase_does_not_warn(self):
        out = run_out(inject(self.T, "<h1>次に進む条件は、1部門で2か月試行して効果を確かめることである</h1>"))
        self.assertNotIn("接続詞で始まる", out)

    def test_word_starting_with_mata_does_not_warn(self):
        out = run_out(inject(self.T, "<h1>またがる2部門で試行し、効果を確かめてから全社へ広げる</h1>"))
        self.assertNotIn("接続詞で始まる", out)


class ProductionMeta(unittest.TestCase):
    """SKILL.md「本スキル使用の注釈」: ツール名は最終ページ（裏表紙）の出典行だけ。他ページは FAIL。"""
    BODY = "承認待ちは平均3日"

    def test_tool_name_on_content_page_fails(self):
        code, fails = run(inject(self.BODY, self.BODY + "<p>本資料は consulting-pptx-skill で作成</p>"))
        self.assertTrue(any("制作メタ" in f for f in fails), fails)

    def test_tool_name_on_back_cover_passes(self):
        code, fails = run(GOOD)   # fixture の裏表紙には出典行に注釈が入っている
        self.assertFalse(any("制作メタ" in f for f in fails), fails)

    def test_credit_sentence_on_content_page_only_warns(self):
        html = inject(self.BODY, self.BODY + "<p>本資料は生成AIツールで作成</p>")
        self.assertFalse(any("制作メタ" in f for f in run(html)[1]))
        self.assertTrue(any("で作成" in w for w in warns(html)))

    def test_disclaimer_with_de_sakusei_does_not_fail(self):
        code, fails = run(inject(self.BODY, self.BODY + "<p>本資料は2026年9月時点の公開情報で作成</p>"))   # 免責文は制作メタではない
        self.assertFalse(any("制作メタ" in f for f in fails), fails)

    def test_github_url_as_source_only_warns(self):
        code, fails = run(inject(self.BODY, self.BODY + "<p>出典: github.com/example/repo</p>"))
        self.assertFalse(any("制作メタ" in f for f in fails), fails)


def _has_playwright():
    import os
    return (ROOT / "node_modules" / "playwright").exists() or bool(os.environ.get("PLAYWRIGHT_MODULE_DIR"))


@unittest.skipUnless(_has_playwright(), "npm run setup で playwright を入れると実行される")
class EmptyArea(unittest.TestCase):
    def layout(self, html):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "deck.html"
            p.write_text(html, encoding="utf8")
            r = subprocess.run(["node", str(ROOT / "scripts" / "check_layout.mjs"), str(p)],
                               capture_output=True, text=True, cwd=ROOT)
        return r.returncode, r.stdout

    def test_good_deck_has_no_empty_page(self):
        code, out = self.layout(GOOD)
        self.assertEqual(code, 0, out)

    def test_half_empty_page_fires(self):
        html = re.sub(r'\s*<tr><td class="ax">2026年7月</td>.*?変わらず</td></tr>', "", GOOD, flags=re.S)
        html = html.replace("\n          <li>1人あたりの申請件数が2倍になる</li>\n          <li>承認の遅れが立替者の不満になる</li>", "")
        self.assertNotEqual(html, GOOD)
        code, out = self.layout(html)
        self.assertEqual(code, 1, out)
        self.assertIn("p3: 版面の", out)


@unittest.skipUnless(_has_playwright(), "npm run setup で playwright を入れると実行される")
class FillMeaning(unittest.TestCase):
    """§5.2 表の塗りには同じ色の凡例チップ／§6 表の見出しは塗らない"""
    ROW = '<tr><td class="ax">試行</td><td>2か月</td><td>営業部</td>'
    HEAD = '<th class="ax" style="width:52mm">段階</th>'
    SRC = '<div class="src">出典：架空データ（説明用）</div>\n  </div>\n  <div class="foot"><b>業務自動化の検討</b><span>4</span></div>'

    def layout(self, html):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "deck.html"
            p.write_text(html, encoding="utf8")
            r = subprocess.run(["node", str(ROOT / "scripts" / "check_layout.mjs"), str(p)],
                               capture_output=True, text=True, cwd=ROOT)
        return r.returncode, r.stdout

    def painted(self):
        return inject(self.ROW, '<tr><td class="ax">試行</td><td style="background:#e8dcc8">2か月</td><td style="background:#e8dcc8">営業部</td>')

    def test_fill_without_legend_warns(self):
        code, out = self.layout(self.painted())
        self.assertIn("凡例チップがない", out)

    def test_fill_with_legend_passes(self):
        html = self.painted()
        chip = '<span style="display:inline-block;width:10px;height:10px;background:#e8dcc8"></span> 試行の対象'
        html = html.replace("<h1>1部門で2か月試行し、効果を確かめてから全社へ広げる</h1>",
                            "<h1>1部門で2か月試行し、効果を確かめてから全社へ広げる</h1><div class=\"legend\">" + chip + "</div>", 1)
        code, out = self.layout(html)
        self.assertNotIn("凡例チップがない", out)

    def test_oklch_fill_without_legend_warns(self):
        html = inject(self.ROW, '<tr><td class="ax">試行</td><td style="background:oklch(0.9 0.04 80)">2か月</td><td>営業部</td>')
        code, out = self.layout(html)
        self.assertIn("凡例チップがない", out)

    def test_oklch_fill_with_legend_passes(self):
        html = inject(self.ROW, '<tr><td class="ax">試行</td><td style="background:oklch(0.9 0.04 80)">2か月</td><td>営業部</td>')
        chip = '<span style="display:inline-block;width:10px;height:10px;background:oklch(0.9 0.04 80)"></span> 試行の対象'
        html = html.replace("<h1>1部門で2か月試行し、効果を確かめてから全社へ広げる</h1>",
                            "<h1>1部門で2か月試行し、効果を確かめてから全社へ広げる</h1><div class=\"legend\">" + chip + "</div>", 1)
        code, out = self.layout(html)
        self.assertNotIn("凡例チップがない", out)

    def test_header_fill_warns(self):
        code, out = self.layout(inject(self.HEAD, '<th class="ax" style="width:52mm;background:#e8dcc8">段階</th>'))
        self.assertIn("表の見出しに塗り", out)


if __name__ == "__main__":
    unittest.main()


class DashAndHeads(unittest.TestCase):
    """§6 非該当「—」は連結ではない／§4.49 基本パーツ集のカラム見出し .colh も対象"""
    def test_label_dash_join_warns(self):
        w = warns(inject("承認待ちは平均3日", "承認待ち — 平均3日"))
        self.assertTrue(any("ダッシュ" in l for l in w), w)

    def test_standalone_na_dash_is_ignored(self):
        w = warns(inject("承認待ちは平均3日", "承認待ちは平均3日</p><p>—</p><p>—（関与しない）"))
        self.assertFalse(any("ダッシュ" in l for l in w), w)

    def test_colh_conjunction_head_warns(self):
        w = warns(inject("承認待ちは平均3日", '承認待ちは平均3日<div class="colh">だから、全社展開を急ぐ</div>'))
        self.assertTrue(any("接続詞" in l for l in w), w)


class HarveyAndSummary(unittest.TestCase):
    """¾ハーベイボールの中心点抜け（FAIL）／エグゼクティブサマリーの羅列（WARN）"""
    BAD_Q3 = "<style>.hb.q3 i{clip-path:polygon(50% 0,100% 0,100% 100%,0 100%,0 50%)}</style>"
    GOOD_Q3 = "<style>.hb.q3 i{clip-path:polygon(50% 50%,50% 0,100% 0,100% 100%,0 100%,0 50%)}</style>"

    def test_q3_without_center_fails(self):
        code, fails = run(GOOD.replace("</head>", self.BAD_Q3 + "</head>", 1))
        self.assertTrue(any("ハーベイボール" in l for l in fails), fails)

    def test_q3_with_center_passes(self):
        code, fails = run(GOOD.replace("</head>", self.GOOD_Q3 + "</head>", 1))
        self.assertFalse(any("ハーベイボール" in l for l in fails), fails)

    FLAT = ('<section class="s"><h1>エグゼクティブサマリー</h1><ul><li>需要はある</li><li>利益は出る</li></ul>'
            '<div class="foot"><span>9</span></div></section>')
    NESTED = ('<section class="s"><h1>エグゼクティブサマリー</h1><ul><li>需要はある<ul><li>調査で確認</li>'
              '<li>価格も通る</li></ul></li></ul><div class="foot"><span>9</span></div></section>')

    def test_flat_summary_warns(self):
        w = warns(GOOD.replace("</body>", self.FLAT + "</body>", 1))
        self.assertTrue(any("エグゼクティブサマリー" in l for l in w), w)

    def test_nested_summary_passes(self):
        w = warns(GOOD.replace("</body>", self.NESTED + "</body>", 1))
        self.assertFalse(any("入れ子のブレット" in l for l in w), w)


def _has_pptx():
    try:
        import pptx  # noqa: F401
        return True
    except ImportError:
        return False


@unittest.skipUnless(_has_pptx(), "python-pptx が入っていると実行される")
class PptxLayout(unittest.TestCase):
    """check_deck_layout.py: 図形に隠れた文字・中身の無い箱（線は箱と取り違えない）"""

    def _deck(self, build):
        from pptx import Presentation
        from pptx.util import Inches
        from pptx.enum.shapes import MSO_SHAPE
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        s = prs.slides.add_slide(prs.slide_layouts[6])
        build(s, Inches, MSO_SHAPE)
        d = tempfile.mkdtemp()
        p = Path(d) / "deck.pptx"
        prs.save(p)
        sys.path.insert(0, str(ROOT / "scripts"))
        import check_deck_layout
        return [k for _, k, _ in check_deck_layout.check(str(p))]

    @staticmethod
    def _solid(shape):
        shape.fill.solid()
        return shape

    def test_text_under_opaque_box_fires(self):
        def build(s, In, M):
            tb = s.shapes.add_textbox(In(1), In(1), In(4), In(1))
            tb.text_frame.text = "隠れている文字"
            self._solid(s.shapes.add_shape(M.RECTANGLE, In(0.9), In(0.9), In(4.2), In(1.2)))
        self.assertIn("hidden", self._deck(build))

    def test_visible_text_and_thin_line_do_not_fire(self):
        def build(s, In, M):
            tb = s.shapes.add_textbox(In(1), In(1), In(4), In(1))
            tb.text_frame.text = "見えている文字"
            self._solid(s.shapes.add_shape(M.RECTANGLE, In(1), In(3), In(1.3), In(0.01)))
        kinds = self._deck(build)
        self.assertNotIn("hidden", kinds)
        self.assertNotIn("empty", kinds)

    def test_empty_filled_box_fires(self):
        def build(s, In, M):
            self._solid(s.shapes.add_shape(M.RECTANGLE, In(1), In(1), In(2), In(1)))
        self.assertIn("empty", self._deck(build))


class MemberFeedbackChecks(unittest.TestCase):
    """メンバーの作業記録から足した検査：文末の「。」／ブレット記号の混在／表のページの連続／操作用の部品"""

    def page(self, n, body, title="検査の見本として置いたページのタイトル"):
        return f'<section class="s"><h1>{title}</h1>{body}<div class="foot"><span>{n}</span></div></section>'

    def test_trailing_period_warns_and_clean_passes(self):
        bad = self.page(9, "<ul><li>添付漏れを提出前に警告する。</li></ul>")
        ok = self.page(9, "<ul><li>添付漏れを提出前に警告する</li></ul>")
        self.assertTrue(any("§7.25" in l for l in warns(GOOD.replace("</body>", bad + "</body>", 1))))
        self.assertFalse(any("§7.25" in l for l in warns(GOOD.replace("</body>", ok + "</body>", 1))))

    def test_mixed_bullet_marks_warn(self):
        bad = self.page(9, "<div>• 添付漏れを警告する</div><div>● 勘定科目を提示する</div>")
        ok = self.page(9, "<div>• 添付漏れを警告する</div><div>• 勘定科目を提示する</div>")
        self.assertTrue(any("§7.3" in l for l in warns(GOOD.replace("</body>", bad + "</body>", 1))))
        self.assertFalse(any("§7.3" in l for l in warns(GOOD.replace("</body>", ok + "</body>", 1))))

    def test_three_table_pages_in_a_row_warn(self):
        t = "<table><tr><th>観点</th><th>内容</th></tr><tr><td>費用</td><td>年額で比べる</td></tr></table>"
        three = "".join(self.page(9 + k, t) for k in range(3))
        two = "".join(self.page(9 + k, t) for k in range(2))
        self.assertTrue(any("§4.96" in l for l in warns(GOOD.replace("</body>", three + "</body>", 1))))
        self.assertFalse(any("§4.96" in l for l in warns(GOOD.replace("</body>", two + "</body>", 1))))

    def test_buttons_must_hide_in_print(self):
        bad = self.page(9, "<button>次へ</button>")
        hidden = bad + "<style>@media print{button{display:none}}</style>"
        self.assertTrue(any("§1" in l for l in warns(GOOD.replace("</body>", bad + "</body>", 1))))
        self.assertFalse(any("§1:" in l for l in warns(GOOD.replace("</body>", hidden + "</body>", 1))))


class WeeklyFeedback20261012(unittest.TestCase):
    """10/12 回収分：改行済みタイトルの字数 WARN／正本から差し込んだ定型ページ／構成変更後に残った P.n"""

    def page(self, n, body, title="検査の見本として置いたページのタイトル", cls="s"):
        return f'<section class="{cls}"><h1>{title}</h1>{body}<div class="foot"><span>{n}</span></div></section>'

    LONG = "提出時の確認と承認経路の見直しで、差し戻しと待ち時間を減らし、承認待ちを当日に縮める"

    def test_long_title_without_break_warns(self):
        w = warns(GOOD.replace("</body>", self.page(9, "<p>本文</p>", self.LONG) + "</body>", 1))
        self.assertTrue(any("p7: タイトル 全角換算" in l for l in w), w)

    def test_title_broken_with_br_does_not_warn(self):
        t = self.LONG.replace("減らし、", "減らし、<br>", 1)
        w = warns(GOOD.replace("</body>", self.page(9, "<p>本文</p>", t) + "</body>", 1))
        self.assertFalse(any("p7: タイトル 全角換算" in l for l in w), w)

    def test_std_page_skips_text_checks(self):
        body = "<p>当社は2020年に創業した。東京に本社を置く。</p>"
        code, fails = run(GOOD.replace("</body>", self.page(9, body) + "</body>", 1))
        self.assertTrue(any("2文以上" in f for f in fails), fails)
        code, fails = run(GOOD.replace("</body>", self.page(9, body, cls="s std") + "</body>", 1))
        self.assertFalse(any("2文以上" in f for f in fails), fails)

    def test_page_ref_beyond_deck_warns(self):
        bad = self.page(9, "<p>詳細は → P.12</p>")
        ok = self.page(9, "<p>詳細は → P.3</p>")
        self.assertTrue(any("P.12" in l for l in warns(GOOD.replace("</body>", bad + "</body>", 1))))
        self.assertFalse(any("ページ数" in l for l in warns(GOOD.replace("</body>", ok + "</body>", 1))))


@unittest.skipUnless(_has_playwright(), "npm run setup で playwright を入れると実行される")
class LayoutFeedback20261012(unittest.TestCase):
    """check_layout の WARN：並べた要素の揃い（§4.51）／片側の空き（§4.54）／ハイライトの数（§4.57）"""

    def layout(self, body, cls="s"):
        page = (f'<section class="{cls}"><h1>検査の見本として置いたページのタイトル</h1>{body}'
                '<div class="foot"><span>9</span></div></section>')
        html = GOOD.replace("</body>", page + "</body>", 1)
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "deck.html"
            p.write_text(html, encoding="utf8")
            r = subprocess.run(["node", str(ROOT / "scripts" / "check_layout.mjs"), str(p)],
                               capture_output=True, text=True, cwd=ROOT, encoding="utf8")
        return [l for l in r.stdout.splitlines() if l.startswith("p7") or "p7:" in l]

    COL = '<div class="cmp"><h3>{h}</h3><ul><li>{t}</li><li>{t}</li></ul></div>'

    def two_cols(self, cols):
        c = "".join(self.COL.format(h=h, t="申請の内容を確認して承認者に回す") for h in ("現状", "導入後"))
        return f'<div style="display:grid;grid-template-columns:{cols};gap:24px;height:380px">{c}</div>'

    def test_unequal_contrast_columns_warn(self):
        out = self.layout(self.two_cols("1fr 1.4fr"))
        self.assertTrue(any("§4.51" in l for l in out), out)

    def test_equal_contrast_columns_pass(self):
        out = self.layout(self.two_cols("1fr 1fr"))
        self.assertFalse(any("§4.51" in l for l in out), out)

    def test_right_side_gap_warns(self):
        body = '<div style="width:50%;height:380px"><ul><li>申請の内容を確認する</li></ul><div style="height:300px;border:1px solid #333"></div></div>'
        self.assertTrue(any("右側" in l for l in self.layout(body)))
        self.assertFalse(any("右側" in l for l in self.layout(self.two_cols("1fr 1fr"))))

    def test_two_highlighted_rows_warn(self):
        row = '<tr><td>費用</td><td>年額で比べる</td></tr>'
        hl = '<tr><td style="background:#eee">費用</td><td style="background:#eee">年額で比べる</td></tr>'
        t = lambda rows: f'<table style="height:380px"><tr><th>観点</th><th>内容</th></tr>{rows}</table>'
        self.assertTrue(any("§4.57" in l for l in self.layout(t(hl + row + hl))))
        self.assertFalse(any("§4.57" in l for l in self.layout(t(hl + row + row))))

    def test_std_page_skips_layout_checks(self):
        body = '<div style="width:50%;height:380px"><ul><li>会社概要</li></ul><div style="height:300px;border:1px solid #333"></div></div>'
        self.assertFalse(any("右側" in l for l in self.layout(body, cls="s std")))


@unittest.skipUnless(_has_playwright(), "npm run setup で playwright を入れると実行される")
class TableStretchAndLegend(LayoutFeedback20261012):
    """check_layout：表を縦に引き伸ばさない（§4.12 FAIL）／凡例はグラフの見出しの上に右揃え（§5.18 WARN）"""

    ROWS = '<tr><th>部門</th><th>件数</th></tr><tr><td>営業部</td><td>月に約120件</td></tr><tr><td>管理部</td><td>月に約40件</td></tr>'
    CHART = '<svg viewBox="0 0 400 200" style="height:70mm"><rect x="20" y="40" width="80" height="160" fill="#5A3921"/><rect x="160" y="100" width="80" height="100" fill="#C5A681"/></svg>'
    LEGEND = '<div class="legend"><span><i style="display:inline-block;width:4mm;height:4mm;background:#5A3921"></i>本社</span><span><i style="display:inline-block;width:4mm;height:4mm;background:#C5A681"></i>拠点</span></div>'

    def test_stretched_table_fails(self):
        out = self.layout(f'<table style="height:420px">{self.ROWS}</table>')
        self.assertTrue(any("引き伸ばされている" in l for l in out), out)

    def test_natural_table_passes(self):
        out = self.layout(f'<table>{self.ROWS}</table>')
        self.assertFalse(any("引き伸ば" in l for l in out), out)

    def test_legend_below_chart_warns(self):
        out = self.layout(f'<div>{self.CHART}{self.LEGEND}</div>')
        self.assertTrue(any("凡例がグラフの下" in l for l in out), out)

    def test_legend_above_chart_passes(self):
        out = self.layout(f'<div>{self.LEGEND}{self.CHART}</div>')
        self.assertFalse(any("凡例がグラフの下" in l for l in out), out)

    def test_row_header_th_table_is_checked(self):
        rows = '<tr><th>部門</th><th>件数</th></tr><tr><th>営業部</th><td>月に約120件</td></tr><tr><th>管理部</th><td>月に約40件</td></tr>'
        out = self.layout(f'<table style="height:420px">{rows}</table>')
        self.assertTrue(any("引き伸ばされている" in l for l in out), out)

    def test_natural_table_with_list_cells_passes(self):
        rows = ('<tr><th>部門</th><th>課題</th></tr>'
                '<tr><td>営業部</td><td><ul><li>入力の二重化</li><li>集計の手作業</li></ul></td></tr>'
                '<tr><td>管理部</td><td><ul><li>承認の待ち</li><li>差し戻し</li></ul></td></tr>')
        out = self.layout(f'<table>{rows}</table>')
        self.assertFalse(any("引き伸ば" in l or "行ぶん高い" in l for l in out), out)

    def test_lg_size_class_is_not_a_legend(self):
        out = self.layout(f'<div>{self.CHART}<div class="text-lg">図の読み方の補足</div></div>')
        self.assertFalse(any("凡例がグラフの下" in l for l in out), out)

    def test_legend_under_single_tile_is_not_below_chart(self):
        tile = '<div style="background:#5A3921;color:#fff;width:60mm;height:30mm">月に約120件の問い合わせ</div>'
        out = self.layout(f'<div>{tile}{self.LEGEND}</div>')
        self.assertFalse(any("凡例がグラフの下" in l for l in out), out)

    def test_legend_directly_in_section_below_chart_warns(self):
        out = self.layout(f'{self.CHART}{self.LEGEND}')
        self.assertTrue(any("凡例がグラフの下" in l for l in out), out)

    # 親クラスのテストを二重に回さない
    test_unequal_contrast_columns_warn = test_equal_contrast_columns_pass = None
    test_right_side_gap_warns = test_two_highlighted_rows_warn = test_std_page_skips_layout_checks = None


@unittest.skipUnless(_has_playwright(), "npm run setup で playwright を入れると実行される")
class ColumnRuleAlignment(LayoutFeedback20261012):
    """check_layout：左右2列の最初の横罫の高さを揃える（§4.21 WARN）"""

    AXH = '<div><div style="height:38px;border-bottom:1.8px solid #333;font-weight:700">工数の推移</div><div style="height:200px;background:#C5A681"></div></div>'
    TBL = ('<table style="border-collapse:collapse"><tr><td style="border-top:1.8px solid #333">営業部</td><td style="border-top:1.8px solid #333">月に約120件</td></tr>'
           '<tr><td>管理部</td><td>月に約40件</td></tr></table>')
    HDR = '<div style="height:38px;border-bottom:1.8px solid #333;font-weight:700">部門ごとの件数</div>'

    def cols(self, left, right):
        return f'<div style="display:grid;grid-template-columns:1fr 1fr;gap:24px"><div>{left}</div><div>{right}</div></div>'

    def test_headerless_table_next_to_axis_title_warns(self):
        out = self.layout(self.cols(self.AXH, self.TBL))
        self.assertTrue(any("最初の横罫" in l for l in out), out)

    def test_table_with_same_height_header_passes(self):
        out = self.layout(self.cols(self.AXH, self.HDR + self.TBL.replace("border-top:1.8px solid #333", "")))
        self.assertFalse(any("最初の横罫" in l for l in out), out)



class BotFeedback20261019(unittest.TestCase):
    """check_deck：文字化け（FAIL）・ページ番号の連番（FAIL）・1項目だけの箇条・本文のですます・締めのページ・keep-all・番号つき項目名（WARN）"""

    def test_garbled_text_fails(self):
        code, fails = run(inject("確認作業が1日1時間を超える", "確認作業が1日1�間を超える"))
        self.assertTrue(any("文字化け" in l for l in fails), fails)

    def test_page_number_gap_fails(self):
        code, fails = run(inject("<span>3</span>", "<span>5</span>"))
        self.assertTrue(any("連番でない" in l for l in fails), fails)
        self.assertFalse(any("連番でない" in l for l in run(GOOD)[1]))

    def test_single_item_list_warns(self):
        html = inject("<li>月末に承認が集中する</li>", "</ul><ul><li>月末に承認が集中する</li></ul><ul>")
        self.assertTrue(any("1つだけの箇条" in l for l in warns(html)))
        self.assertFalse(any("1つだけの箇条" in l for l in warns(GOOD)))

    def test_body_desu_masu_warns(self):
        html = inject("確認作業が1日1時間を超える", "確認作業が1日1時間を超えます")
        self.assertTrue(any("本文がですます調" in l for l in warns(html)))
        self.assertFalse(any("本文がですます調" in l for l in warns(GOOD)))

    def test_closing_slide_is_not_desu_masu_fail(self):
        last = GOOD.rfind('<section class="s cover">')
        html = GOOD[:last] + GOOD[last:].replace("<section class=\"s cover\">", "<section class=\"s cover\"><h1>ありがとうございました</h1>", 1)
        code, fails = run(html)
        self.assertFalse(any("ですます" in l for l in fails), fails)

    def test_keep_all_warns(self):
        html = inject("</style>", "h1{word-break:keep-all}</style>")
        self.assertTrue(any("keep-all" in l for l in warns(html)))

    def test_numbered_labels_mismatch_warns(self):
        html = inject("確認作業が1日1時間を超える", "①提出前の確認").replace("月末に承認が集中する", "①事前チェック", 1)
        html = html.replace("<span>3</span>", "<span>3</span><i>①事前チェック</i>", 1)
        self.assertTrue(any("番号つきの項目名" in l for l in warns(html)))

    def test_mixed_parts_header_label_passes(self):
        # 基本パーツ（.bar/.date）と追加パーツ（.kicker → new_deck で .hdr/.date）の混在は正規のデッキ
        with tempfile.TemporaryDirectory() as d:
            out = Path(d) / "mixed.html"
            subprocess.run([sys.executable, str(ROOT / "scripts" / "new_deck.py"),
                            "--parts", "b01,b02,b05,m05,b06,m10,b09,m20,b10", "-o", str(out)],
                           capture_output=True, check=True)
            r = subprocess.run([sys.executable, str(CHECK), str(out), *NO_BRAND], capture_output=True, text=True, encoding="utf8")
        self.assertFalse(any("右上のラベル" in l for l in r.stdout.splitlines()), r.stdout[-600:])
        self.assertFalse(any("右上のラベル" in l for l in warns(GOOD)))

    def test_new_deck_fits_narrow_window(self):
        # 画面幅がスライド幅より狭いと左右が切れていた。生成物に等比縮小の script と、印刷時に縮小を外す指定があること
        with tempfile.TemporaryDirectory() as d:
            out = Path(d) / "fit.html"
            subprocess.run([sys.executable, str(ROOT / "scripts" / "new_deck.py"),
                            "--parts", "b01,b02,m05,b10", "-o", str(out)], capture_output=True, check=True)
            html = out.read_text(encoding="utf8")
        self.assertIn('wrap.className = "fitwrap"', html)
        self.assertRegex(html, r"@media print\{\s*\.fitwrap\{width:auto!important")  # .s の内側にスコープされていない

    def test_missing_header_label_warns(self):
        html = inject('<div class="date">打ち手</div>', '<div class="date"></div>')
        self.assertTrue(any("右上のラベル" in l for l in warns(html)), warns(html))


@unittest.skipUnless(_has_playwright(), "npm run setup で playwright を入れると実行される")
class LayoutBotFeedback20261019(LayoutFeedback20261012):
    """check_layout：箱からのはみ出し（FAIL）・箱の中の空き・下の空き30%・同じ作りで違う大きさ・文の途中の太字・混在要素の文字の重なり（WARN）"""

    def test_text_overflowing_box_fails(self):
        body = '<div style="width:120px;height:40px;overflow:hidden;background:#EFEEE8;white-space:nowrap">申請の内容を確認して承認者に回す処理を自動化する</div><div style="height:300px;border:1px solid #333">図</div>'
        self.assertTrue(any("箱の外に" in l for l in self.layout(body)))

    def test_text_inside_box_passes(self):
        body = '<div style="width:600px;background:#EFEEE8">申請の内容を確認して承認者に回す</div><div style="height:300px;border:1px solid #333">図</div>'
        self.assertFalse(any("箱の外に" in l for l in self.layout(body)))

    def test_empty_box_warns(self):
        body = '<div style="height:380px;background:#EFEEE8;padding:8px"><p>申請の内容を確認する</p></div>'
        self.assertTrue(any("箱の中が" in l for l in self.layout(body)))

    def test_equal_height_sibling_boxes_pass(self):
        full = "".join(f"<p>申請の内容を確認する手順その{i}</p>" for i in range(9))
        body = (f'<div style="display:grid;grid-template-columns:1fr 1fr;gap:16px;height:380px">'
                f'<div class="cd" style="background:#EFEEE8">{full}</div><div class="cd" style="background:#EFEEE8"><p>申請の内容を確認する</p></div></div>')
        self.assertFalse(any("箱の中が" in l for l in self.layout(body)))

    def test_bottom_space_over_30_warns(self):
        body = '<div style="height:370px;border:1px solid #333"><p>申請の内容を確認する</p></div>'
        out = self.layout(body)
        self.assertTrue(any("30%超" in l for l in out), out)

    def test_same_role_different_size_warns(self):
        body = ('<ul><li style="font-size:14px">申請の内容を確認する</li><li style="font-size:18px">承認者に回して結果を返す</li></ul>'
                '<div style="height:300px;border:1px solid #333">図</div>')
        self.assertTrue(any("違う大きさ" in l for l in self.layout(body)))

    def test_mid_sentence_bold_warns(self):
        li = "<li>申請の<b>内容</b>を確認し<b>承認者</b>に回して<b>結果</b>を返す</li>"
        body = f'<ul>{li}<li>月末に承認が集中する</li></ul><div style="height:300px;border:1px solid #333">図</div>'
        self.assertTrue(any("途中の太字" in l for l in self.layout(body)))

    def test_overlap_in_mixed_element_warns(self):
        body = ('<div style="position:relative;height:340px;border:1px solid #333"><p style="position:absolute;top:10px;left:10px">申請の<b>内容</b>を確認する</p>'
                '<p style="position:absolute;top:12px;left:14px">承認者に回して結果を返す</p></div>')
        self.assertTrue(any("文字が重なっている" in l for l in self.layout(body)))


def _has_chrome():
    import os
    import shutil
    paths = [os.environ.get("CHROME_PATH") or "", "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
             "C:/Program Files/Google/Chrome/Application/chrome.exe", "/usr/bin/google-chrome", "/usr/bin/chromium"]
    return bool(shutil.which("node")) and any(p and Path(p).exists() for p in paths)


@unittest.skipUnless(_has_pptx() and _has_chrome(), "python-pptx・Node・Chrome が要る")
class PptxCssZoom(unittest.TestCase):
    """html_to_pptx：CSS の zoom で縮めたスライド（追加パーツ集の .slide）の文字・行間・罫線を、縮小後の大きさで書き出す"""

    HTML = ('<!doctype html><html lang="ja"><head><meta charset="utf-8"><style>'
            '@page{size:338.67mm 190.5mm;margin:0}body{margin:0}'
            '.s{width:338.67mm;height:190.5mm;position:relative;page-break-after:always;background:#fff}'
            '.slide{width:1600px;height:900px;zoom:0.8;position:relative;page-break-after:always;background:#fff}'
            'h1{position:absolute;left:40px;top:40px;margin:0;font-size:40px;line-height:50px;border-bottom:5px solid #000}'
            '</style></head><body>'
            '<section class="s"><h1>承認の手順を見直す</h1></section>'
            '<section class="slide"><h1>承認の手順を見直す</h1></section>'
            '</body></html>')

    def convert(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "deck.html"
            p.write_text(self.HTML, encoding="utf8")
            out = Path(d) / "deck.pptx"
            subprocess.run([sys.executable, str(ROOT / "scripts" / "html_to_pptx.py"), str(p), str(out)],
                           capture_output=True, text=True, check=True)
            got = []
            for sl in Presentation(str(out)).slides:
                text = next(sh for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text)
                line = next(sh for sh in sl.shapes if not (sh.has_text_frame and sh.text_frame.text))
                para = text.text_frame.paragraphs[0]
                got.append((para.runs[0].font.size.pt, para.line_spacing.pt, line.line.width.pt))
            return got

    def test_zoomed_slide_uses_rendered_sizes(self):
        plain, zoomed = self.convert()
        # 1px = 0.75pt。縮小のないスライドは 40px→30pt、zoom:0.8 のスライドは 32px→24pt
        self.assertAlmostEqual(plain[0], 30.0, delta=0.5)
        self.assertAlmostEqual(zoomed[0], 24.0, delta=0.5)
        self.assertAlmostEqual(zoomed[1], 30.0, delta=0.5)   # 行間 50px×0.8
        self.assertAlmostEqual(zoomed[2], 3.0, delta=0.1)    # 罫線 5px×0.8


@unittest.skipUnless(_has_pptx() and _has_playwright(), "python-pptx と playwright が要る")
class PptxVideoAndExternal(unittest.TestCase):
    """html_to_pptx：<video> を動画として埋め込む／--external で section.internal を除く"""

    def test_video_embedded_and_internal_skipped(self):
        import shutil
        if not shutil.which("ffmpeg"):
            self.skipTest("ffmpeg が要る")
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as d:
            mp4 = Path(d) / "clip.mp4"
            subprocess.run(["ffmpeg", "-loglevel", "error", "-f", "lavfi", "-i", "color=c=gray:s=320x180:d=1", "-pix_fmt", "yuv420p", str(mp4)], check=True)
            extra = ('<section class="s"><h1>画面の操作を動画で示す</h1><video src="clip.mp4" style="width:320px;height:180px"></video>'
                     '<div class="foot"><span>9</span></div></section>'
                     '<section class="s internal"><h1>社内用の確認事項を残しておく</h1><p>社内用</p></section>')
            p = Path(d) / "deck.html"
            p.write_text(GOOD.replace("</body>", extra + "</body>", 1), encoding="utf8")
            out = Path(d) / "deck.pptx"
            subprocess.run([sys.executable, str(ROOT / "scripts" / "html_to_pptx.py"), str(p), str(out), "--external"],
                           capture_output=True, text=True, check=True)
            prs = Presentation(str(out))
            n_html = GOOD.count("<section")
            self.assertEqual(len(prs.slides), n_html + 1)
            xml = "".join(sl.part.blob.decode("utf8", "ignore") for sl in prs.slides)
            self.assertIn("p:videoFile", xml.replace("a:videoFile", "p:videoFile"))


class TypeIdLeftover(unittest.TestCase):
    """check_deck：パーツの英語の型ID（「05  chart_insight」）が残っていたら FAIL"""

    def test_type_id_fails(self):
        code, fails = run(inject("<span>3</span>", "<span>3</span><i>05  chart_insight</i>"))
        self.assertTrue(any("型ID" in l for l in fails), fails)

    def test_good_deck_has_no_type_id(self):
        self.assertFalse(any("型ID" in l for l in run(GOOD)[1]))

    def test_table_number_and_snake_case_is_not_type_id(self):
        # 表の「05」「user_id」のような並びは型IDではない
        code, fails = run(inject("<span>3</span>", "<span>3</span><table><tr><td>05</td><td>user_id</td></tr></table>"))
        self.assertFalse(any("型ID" in l for l in fails), fails)

    def test_single_word_type_id_fails(self):
        for tid in ("34  gantt", "15  matrix_2x2"):
            code, fails = run(inject("<span>3</span>", f"<span>3</span><i>{tid}</i>"))
            self.assertTrue(any("型ID" in l for l in fails), (tid, fails))


@unittest.skipUnless(_has_playwright(), "npm run setup で playwright を入れると実行される")
class BarAxisAndColors(LayoutFeedback20261012):
    """check_layout：棒グラフの項目名が軸の線より上（§5.26）・下端のずれ・凡例なしの色分け（§5.5）"""

    def chart(self, labels_below=True, colors=("#322014",) * 4, legend=""):
        wrap = "position:relative;display:flex;flex-direction:column;justify-content:flex-end;align-items:center;height:100%"
        lab = "position:absolute;top:calc(100% + 8px);left:0;right:0;text-align:center" if labels_below else "margin-top:8px;text-align:center"
        bars = "".join(f'<div style="{wrap}"><div>{v}</div><div style="width:70%;height:{v}%;background:{c}"></div><div style="{lab}">部門{i}</div></div>'
                       for i, (v, c) in enumerate(zip((30, 50, 70, 90), colors)))
        return (f'{legend}<div style="display:grid;grid-template-columns:repeat(4,1fr);gap:20px;height:300px;'
                f'border-bottom:2px solid #333;margin-bottom:52px">{bars}</div>')

    def test_labels_above_axis_warn(self):
        out = self.layout(self.chart(labels_below=False))
        self.assertTrue(any("軸の線より上" in l for l in out), out)

    def test_labels_below_axis_pass(self):
        out = self.layout(self.chart())
        self.assertFalse(any("軸の線より上" in l or "下端が" in l for l in out), out)

    def test_colors_without_legend_warn(self):
        out = self.layout(self.chart(colors=("#322014", "#C5A681", "#5A3921", "#322014")))
        self.assertTrue(any("凡例が無い" in l for l in out), out)

    def test_colors_with_legend_pass(self):
        out = self.layout(self.chart(colors=("#322014", "#C5A681", "#5A3921", "#322014"), legend='<div class="legend">凡例</div>'))
        self.assertFalse(any("凡例が無い" in l for l in out), out)

    def test_single_accent_bar_pass(self):
        # 目を向けたい1本だけ色を変えるのは推奨どおり（§5.26/§5.5）
        out = self.layout(self.chart(colors=("#322014", "#322014", "#322014", "#C5A681")))
        self.assertFalse(any("凡例が無い" in l for l in out), out)

    def test_card_below_chart_is_not_axis(self):
        # 棒の下に置いた下罫つきのカードを軸と取り違えない
        out = self.layout(self.chart().replace("border-bottom:2px solid #333;margin-bottom:52px", "margin-bottom:0")
                          + '<div style="border-bottom:1px solid #ccc;padding:40px 8px 8px">補足のカード</div>')
        self.assertFalse(any("軸の線より上" in l for l in out), out)

    def test_two_charts_side_by_side_checked_separately(self):
        # 2カラムに並んだ2つの棒グラフを1つに合流させない（棒を器の直下に並べる作り）
        def bars(hs, colors):
            return "".join(f'<div style="width:50px;height:{h}px;background:{c}"></div>' for h, c in zip(hs, colors))
        one = "display:flex;align-items:flex-end;gap:16px;height:260px;border-bottom:2px solid #333"
        body = (f'<div style="display:flex;gap:80px"><div style="{one}">{bars((80, 140, 220), ("#322014", "#322014", "#C5A681"))}</div>'
                f'<div style="{one}">{bars((60, 120, 180, 200), ("#322014",) * 4)}</div></div>')
        out = self.layout(body)
        self.assertFalse(any("凡例が無い" in l or "下端が" in l for l in out), out)


@unittest.skipUnless(_has_playwright(), "npm run setup で playwright を入れると実行される")
class RowSlackAndHeadingSize(LayoutFeedback20261012):
    """check_layout：表の行の間延び（§4.12）・見出しが下の文字より小さい（§6・§4.21）"""

    ROWS = "".join(f'<tr><td style="padding:{{p}};line-height:{{lh}}">部門{i}</td><td style="padding:{{p}};line-height:{{lh}}">月に約{i}0件を処理する</td></tr>' for i in range(1, 4))

    def test_loose_rows_warn(self):
        out = self.layout("<table>" + self.ROWS.format(p="5mm 2mm", lh="1.7") + "</table>")
        self.assertTrue(any("間延び" in l for l in out), out)

    def test_tight_rows_pass(self):
        out = self.layout("<table>" + self.ROWS.format(p="2.5mm 2mm", lh="1.4") + "</table>")
        self.assertFalse(any("間延び" in l for l in out), out)

    def test_small_heading_warns(self):
        body = '<h3 style="font-size:12px">部門ごとの件数</h3><div><p style="font-size:18px">営業部は月に約120件を処理する</p></div>'
        self.assertTrue(any("その下の文字より小さい" in l for l in self.layout(body)))

    def test_big_heading_passes(self):
        body = '<h3 style="font-size:20px">部門ごとの件数</h3><div><p style="font-size:15px">営業部は月に約120件を処理する</p><p style="font-size:30px">72%</p></div>'
        self.assertFalse(any("その下の文字より小さい" in l for l in self.layout(body)))


class StructureF(unittest.TestCase):
    """要点版と正本のずれ・出典の照合・check_all の動作"""

    def test_digest_refs_exist_in_canon(self):
        digest = (ROOT / "references" / "rules-digest.md").read_text(encoding="utf8")
        refs = sorted(set(re.findall(r"§(\d+\.\d+)", digest)))
        self.assertTrue(refs)
        r = subprocess.run([sys.executable, str(ROOT / "scripts" / "rules.py"), "--sec", *refs], capture_output=True, text=True)
        self.assertNotIn("見つからない", r.stderr, r.stderr)

    def test_digest_has_all_phases(self):
        digest = (ROOT / "references" / "rules-digest.md").read_text(encoding="utf8")
        for ph in ("plan", "build", "chart", "table", "process", "text", "check"):
            self.assertIn(f"## {ph}", digest)

    def test_sources_missing_number_warns(self):
        with tempfile.TemporaryDirectory() as d:
            src = Path(d) / "sources.md"
            src.write_text("# 出典\n- 月に約120件 社内集計\n", encoding="utf8")
            deck = Path(d) / "deck.html"
            deck.write_text(GOOD, encoding="utf8")
            r = subprocess.run([sys.executable, str(CHECK), str(deck), *NO_BRAND, "--sources", str(src)], capture_output=True, text=True)
        self.assertIn("出典の控えに無い数字", r.stdout)

    def test_sources_match_is_not_substring(self):
        # 控えの「2025」が本文の「5万人」「20倍」を通さない（数字は前後が数字でない境界で照合する）
        with tempfile.TemporaryDirectory() as d:
            src = Path(d) / "sources.md"
            src.write_text("# 出典\n- 2025年 社内集計\n", encoding="utf8")
            deck = Path(d) / "deck.html"
            deck.write_text(inject("差し戻しの削減と承認の自動化で、承認待ちを当日に縮める", "利用者は5万人、処理は20倍に増える"), encoding="utf8")
            r = subprocess.run([sys.executable, str(CHECK), str(deck), *NO_BRAND, "--sources", str(src)], capture_output=True, text=True)
        line = next((l for l in r.stdout.splitlines() if "出典の控えに無い数字" in l), "")
        self.assertIn("「5万」", line)
        self.assertIn("「20倍」", line)

    def test_sources_match_with_thousands_separator(self):
        with tempfile.TemporaryDirectory() as d:
            src = Path(d) / "sources.md"
            src.write_text("# 出典\n- 12,000件 社内集計\n", encoding="utf8")
            deck = Path(d) / "deck.html"
            deck.write_text(inject("差し戻しの削減と承認の自動化で、承認待ちを当日に縮める", "申請は年12000件に達する"), encoding="utf8")
            r = subprocess.run([sys.executable, str(CHECK), str(deck), *NO_BRAND, "--sources", str(src)], capture_output=True, text=True)
        self.assertNotIn("「12000件」", r.stdout)

    def test_check_all_fails_when_a_checker_crashes(self):
        # check_deck が途中で落ちたら（ここでは無いファイル）、0 FAIL として通さず非0で終わる
        with tempfile.TemporaryDirectory() as d:
            r = subprocess.run([sys.executable, str(ROOT / "scripts" / "check_all.py"), str(Path(d) / "nope.html")],
                               capture_output=True, text=True, cwd=d)
        self.assertEqual(r.returncode, 2, r.stdout + r.stderr)
        self.assertIn("最後まで動いていない", r.stderr)
        self.assertNotIn("0 FAIL", r.stdout)

    def test_sources_file_required(self):
        with tempfile.TemporaryDirectory() as d:
            deck = Path(d) / "deck.html"
            deck.write_text(GOOD, encoding="utf8")
            r = subprocess.run([sys.executable, str(CHECK), str(deck), *NO_BRAND, "--sources", str(Path(d) / "nope.md")], capture_output=True, text=True)
        self.assertIn("出典の控え", r.stdout)
        self.assertEqual(r.returncode, 1)

    @unittest.skipUnless(_has_playwright(), "playwright が要る")
    def test_check_all_groups_by_page(self):
        with tempfile.TemporaryDirectory() as d:
            deck = Path(d) / "deck.html"
            deck.write_text(GOOD, encoding="utf8")
            r = subprocess.run([sys.executable, str(ROOT / "scripts" / "check_all.py"), str(deck)], capture_output=True, text=True, cwd=d)
        self.assertIn("FAIL /", r.stdout)
        self.assertIn("== タイトル一覧", r.stdout)


@unittest.skipUnless(_has_playwright(), "npm run setup で playwright を入れると実行される")
class LayoutFeedback20261020(LayoutFeedback20261012):
    """check_layout：文字と罫線の重なり・同じ作りの箱の左端のずれ・行のリストの間延び・左右の行のずれ（WARN）"""

    def test_text_over_rule_warns(self):
        body = ('<div style="height:40px;border-bottom:2px solid #333">指標名、単位</div>'
                '<div style="margin-top:-18px;font-size:20px;font-weight:700">48.5</div><div style="height:300px;border:1px solid #333">図</div>')
        self.assertTrue(any("罫線に重なっている" in l for l in self.layout(body)))

    def test_text_clear_of_rule_passes(self):
        body = '<div style="height:40px;border-bottom:2px solid #333">指標名、単位</div><div style="margin-top:12px;font-size:20px">48.5</div><div style="height:300px;border:1px solid #333">図</div>'
        self.assertFalse(any("罫線に重なっている" in l for l in self.layout(body)))

    def test_box_left_edge_misaligned_warns(self):
        box = '<div class="lf" style="border:1px solid #999;padding:6px;margin-left:{m}px;width:300px">申請の内容を確かめる</div>'
        body = "".join(box.format(m=m) for m in (100, 130, 112)) + '<div style="height:250px;border:1px solid #333">図</div>'
        self.assertTrue(any("左端が" in l for l in self.layout(body)))

    def test_box_left_edge_aligned_passes(self):
        box = '<div class="lf" style="border:1px solid #999;padding:6px;margin-left:100px;width:300px">申請の内容を確かめる</div>'
        body = box * 3 + '<div style="height:250px;border:1px solid #333">図</div>'
        self.assertFalse(any("左端が" in l for l in self.layout(body)))

    ROW = '<div style="border-bottom:1px solid #ccc;padding:{p};display:flex;align-items:center">部門{i}は月に約{i}0件を処理する</div>'

    def test_stretched_row_list_warns(self):
        rows = "".join(self.ROW.format(p="9mm 0", i=i) for i in range(1, 5))
        self.assertTrue(any("行のリストの上下が間延び" in l for l in self.layout(f"<div>{rows}</div>")))

    def test_tight_row_list_passes(self):
        rows = "".join(self.ROW.format(p="3mm 0", i=i) for i in range(1, 5))
        self.assertFalse(any("行のリストの上下が間延び" in l for l in self.layout(f"<div>{rows}</div>")))

    def test_rows_misaligned_across_columns_warns(self):
        left = "".join(self.ROW.format(p="3mm 0", i=i) for i in range(1, 4))
        right = "".join(self.ROW.format(p="3mm 0" if i != 1 else "8mm 0", i=i) for i in range(1, 4))
        body = f'<div style="display:grid;grid-template-columns:1fr 1fr;gap:20px"><div>{left}</div><div>{right}</div></div>'
        self.assertTrue(any("行の位置が" in l for l in self.layout(body)))


class MetaTitle(unittest.TestCase):
    """check_deck：作業の説明のタイトル（「〜を3つに分ける」「〜を示す」「〜を整理する」は WARN、「〜ごとに…を整理する」「〜を一覧にする」は FAIL）"""

    def test_meta_title_warns(self):
        html = inject("差し戻しの削減と承認の自動化で、承認待ちを当日に縮める", "承認が遅れる理由を3つに分ける")
        self.assertTrue(any("作業の説明" in l for l in warns(html)))

    def test_claim_title_passes(self):
        self.assertFalse(any("作業の説明" in l for l in warns(GOOD)))

    def test_seiri_ending_fails(self):
        html = inject("差し戻しの削減と承認の自動化で、承認待ちを当日に縮める", "打ち手ごとに担当・時期・効果を整理する")
        self.assertTrue(any("作業の説明で終わっている" in l for l in run(html)[1]))

    def test_action_with_seiri_ending_is_warn_not_fail(self):
        # 打ち手の言い切り（「承認フローを整理する」「重複した窓口をまとめる」）は FAIL にしない
        for t in ("承認フローを整理する", "重複した窓口をまとめる"):
            html = inject("差し戻しの削減と承認の自動化で、承認待ちを当日に縮める", t)
            self.assertFalse(any("作業の説明" in l for l in run(html)[1]), t)
            self.assertTrue(any("で終わっている" in l for l in warns(html)), t)

    def test_shimesu_ending_warns(self):
        html = inject("差し戻しの削減と承認の自動化で、承認待ちを当日に縮める", "承認待ちの日数の推移を示す")
        self.assertTrue(any("作業の説明で終わっている" in l for l in warns(html)))


class WorkFilesInDeck(unittest.TestCase):
    """check_deck：デッキに作り手の控え（sources.md など）のファイル名が出ていたら FAIL（§2.23）"""

    def test_sources_md_in_note_fails(self):
        html = GOOD.replace("出典：架空データ（説明用）", "出典：架空データ（説明用）。試算の前提は sources.md 参照", 1)
        self.assertTrue(any("作り手の控え" in l for l in run(html)[1]))

    def test_sources_md_after_japanese_fails(self):
        html = GOOD.replace("出典：架空データ（説明用）", "出典：架空データ（算式は添付のsources.mdを参照）", 1)
        self.assertTrue(any("作り手の控え" in l for l in run(html)[1]))

    def test_good_deck_passes(self):
        self.assertFalse(any("作り手の控え" in l for l in run(GOOD)[1]))


class FrameworkLabelsAndCommas(unittest.TestCase):
    """check_deck：思考の枠組みの名前の見出し（FAIL）、日本語の文の中の半角カンマ（WARN）（§7.32）"""

    def test_framework_label_fails(self):
        html = GOOD.replace('<div class="l">現状</div>', '<div class="l">Situation</div>', 1)
        self.assertTrue(any("思考の枠組みの名前" in l for l in run(html)[1]))

    def test_half_comma_warns(self):
        html = GOOD.replace("差し戻しの6割は添付漏れ", "差し戻しの6割は添付漏れ,残りは記入誤り", 1)
        self.assertTrue(any("半角カンマ" in l for l in warns(html)))

    def test_good_deck_passes(self):
        self.assertFalse(any("思考の枠組みの名前" in l or "半角カンマ" in l for l in run(GOOD)[1] + warns(GOOD)))


class DashAndPlaceholders(unittest.TestCase):
    """check_deck：「—」で語句をつなぐ（FAIL。§7.1）、XXXX などの伏せ字（FAIL）"""

    def test_dash_link_fails(self):
        html = inject("差し戻しの削減と承認の自動化で、承認待ちを当日に縮める", "業務自動化 — 承認待ちを当日に縮める")
        self.assertTrue(any("「—」で語句をつないでいる" in l for l in run(html)[1]))

    def test_dash_link_does_not_depend_on_spaces(self):
        # 空白の数やタグの組み方で結果が変わらない: 文中のつなぎは空白がいくつでも FAIL、セル単独の「—」は空白が重なっても可
        for body in ("<p>前年比  —  横ばい</p>", "<p>前年比<span> — </span>横ばい</p>"):
            html = GOOD.replace("差し戻しの6割は添付漏れ", "差し戻しの6割は添付漏れ" + body, 1)
            self.assertTrue(any("「—」で語句をつないでいる" in l for l in run(html)[1]), body)
        html = GOOD.replace("差し戻しの6割は添付漏れ", "差し戻しの6割は添付漏れ<table><tr><td>前年比</td><td>   —   </td></tr></table>", 1)
        self.assertFalse(any("「—」で語句をつないでいる" in l for l in run(html)[1]))

    def test_placeholder_xxxx_fails(self):
        html = GOOD.replace("差し戻しの6割は添付漏れ", "差し戻しの6割は添付漏れ（内線 XXXX）", 1)
        self.assertTrue(any("プレースホルダー" in l for l in run(html)[1]))

    def test_good_deck_passes(self):
        self.assertFalse(any("「—」" in l or "プレースホルダー" in l for l in run(GOOD)[1]))


class CoverPlaceholder(unittest.TestCase):
    """check_deck：表紙・裏表紙に見本の文言（当社・会社名・連絡先・「表紙」のラベル）が残っていたら FAIL"""

    def test_back_cover_placeholder_fails(self):
        last = GOOD.rfind('<section class="s cover">')
        html = GOOD[:last] + GOOD[last:].replace('<section class="s cover">', '<section class="s cover"><div>当社</div>', 1)
        self.assertTrue(any("見本の文言" in l for l in run(html)[1]))

    def test_cover_type_label_fails(self):
        html = GOOD.replace('<section class="s cover">', '<section class="s cover"><div>表紙</div>', 1)
        self.assertTrue(any("見本の文言が残っている「表紙」" in l for l in run(html)[1]))

    def test_good_deck_cover_passes(self):
        self.assertFalse(any("見本の文言" in l for l in run(GOOD)[1]))

    def test_last_content_page_is_not_back_cover(self):
        # 裏表紙の無いデッキの最終ページの比較表「当社 / 競合A」は見本の文言ではない
        page = ('<section class="s"><h1>当社は導入の速さで競合を上回る</h1><table><tr><th>当社</th><th>競合A</th></tr>'
                '<tr><td>2週間</td><td>3か月</td></tr></table><div class="foot"><span>9</span></div></section>')
        last = GOOD.rfind('<section class="s cover">')
        html = GOOD[:last] + page + "</body></html>"
        self.assertFalse(any("見本の文言" in l for l in run(html)[1]))


@unittest.skipUnless(_has_playwright(), "npm run setup で playwright を入れると実行される")
class LayoutReview20261021(LayoutFeedback20261012):
    """check_layout：レビューの指摘（PR #39）の再現"""

    def run_layout(self, html):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "deck.html"
            p.write_text(html, encoding="utf8")
            r = subprocess.run(["node", str(ROOT / "scripts" / "check_layout.mjs"), str(p)],
                               capture_output=True, text=True, cwd=ROOT, encoding="utf8")
        return r.stdout

    def test_label_and_bar_list_is_not_stretched(self):
        # ラベル＋塗りのバー（文字なし div）の縦並びは、バーの高さを空きに数えない
        row = '<div><div style="font-size:12px">{t}</div><div style="height:40px;background:#8a5a3c"></div></div>'
        body = '<div style="width:600px">' + "".join(row.format(t=t) for t in ("東日本", "西日本", "九州")) + "</div>"
        self.assertFalse(any("並べた項目の間" in l for l in self.layout(body)))
        flat = ('<div style="width:600px">' + "".join(f'<div style="font-size:12px">{t}</div><div style="height:40px;background:#8a5a3c"></div>'
                                                      for t in ("東日本", "西日本", "九州")) + "</div>")
        self.assertFalse(any("並べた項目の間" in l for l in self.layout(flat)))

    def test_white_rules_on_dark_background_are_compared(self):
        # 濃色の背景の白い区切り線の太さの不揃いは検出する（白を一律に外さない）
        body = ('<div style="background:#3b2a20;padding:20px;height:300px">'
                '<div class="sep" style="border-top:1px solid #fff;color:#fff">申請</div>'
                '<div class="sep" style="border-top:3px solid #fff;color:#fff">承認</div></div>')
        out = self.run_layout(GOOD.replace("</body>", f'<section class="s"><h1>検査の見本として置いたページのタイトル</h1>{body}'
                                                       '<div class="foot"><span>9</span></div></section></body>', 1))
        self.assertIn("div.sep:top", out)

    def test_zoom_hint_checks_right_edge(self):
        # 列幅を mm で固定した表は、拡大すると右端を越える。文字の下端だけで「はみ出しなし」と勧めず、拡大では直らないと言う
        body = '<div class="c"><table style="table-layout:fixed;width:310mm"><tr><td>申請</td><td>承認</td></tr></table></div>'
        line = next((l for l in self.layout(body) if "が空いている" in l), "")
        self.assertIn("1.05倍にするだけで版面からはみ出す", line)
        self.assertNotIn("zoom:", line)

    def test_div_slide_deck_pdf_pages_are_compared(self):
        # <div class="slide"> のデッキでも PDF のページ数の照合が動く（メイン検査と同じセレクタ）
        page = ('<div class="slide" style="width:338.67mm;height:190.5mm;page-break-after:always;overflow:hidden">'
                '<h1>{t}</h1><p>本文</p></div>')
        html = ('<!doctype html><html lang="ja"><head><meta charset="utf-8"><style>@page{size:338.67mm 190.5mm;margin:0}'
                'body{margin:0}</style></head><body>' + page.format(t="一枚目のタイトル") + page.format(t="二枚目のタイトル")
                + '<div class="slide" style="width:338.67mm;height:400mm"><h1>高すぎる三枚目</h1></div></body></html>')
        self.assertIn("PDF が", self.run_layout(html))


@unittest.skipUnless(_has_playwright(), "npm run setup で playwright を入れると実行される")
class UnevenWrap(LayoutFeedback20261012):
    """check_layout：並べた見出しのうち一部だけが折れている（WARN）"""

    def cols(self, titles):
        cell = '<div><h3 style="font-size:20px;font-weight:700;margin:0">{t}</h3><p>説明の文</p></div>'
        return ('<div style="display:grid;grid-template-columns:repeat(4,160px);gap:20px">'
                + "".join(cell.format(t=t) for t in titles) + '</div><div style="height:260px;border:1px solid #333">図</div>')

    def test_one_heading_wrapped_warns(self):
        out = self.layout(self.cols(["土台づくり", "試行", "本番化", "成果検証と全社への拡大判断"]))
        self.assertTrue(any("折れている" in l for l in out), out)

    def test_all_one_line_passes(self):
        out = self.layout(self.cols(["土台づくり", "試行", "本番化", "拡大判断"]))
        self.assertFalse(any("折れている" in l for l in out), out)


CHROME_APP = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"


class HeadlessShell(unittest.TestCase):
    """PDF 化は Chromium headless shell で行い、Chrome 本体と同じ出力になること"""

    def test_chrome_path_overrides(self):
        import os
        r = subprocess.run([sys.executable, str(ROOT / "scripts" / "find_browser.py")], capture_output=True, text=True,
                           env={**os.environ, "CHROME_PATH": "/opt/custom/chrome"})
        self.assertEqual(r.stdout.strip(), "/opt/custom/chrome")

    @unittest.skipUnless(Path(CHROME_APP).exists(), "Chrome 本体が要る（比較の基準）")
    def test_headless_shell_pdf_matches_chrome(self):
        import shutil
        sys.path.insert(0, str(ROOT / "scripts"))
        from find_browser import headless_shell
        hs = headless_shell()
        if not hs or not shutil.which("pdftoppm"):
            self.skipTest("headless shell（npm run setup）と pdftoppm が要る")
        from PIL import Image, ImageChops
        with tempfile.TemporaryDirectory() as d:
            deck = Path(d) / "deck.html"
            deck.write_text(GOOD, encoding="utf8")
            outs = {}
            for name, exe in (("hs", hs), ("chrome", CHROME_APP)):
                pdf = Path(d) / f"{name}.pdf"
                subprocess.run([exe, "--headless", "--disable-gpu", "--no-pdf-header-footer", f"--print-to-pdf={pdf}", str(deck)],
                               capture_output=True, check=True, timeout=120)
                subprocess.run(["pdftoppm", "-r", "40", "-png", str(pdf), str(Path(d) / name)], check=True)
                outs[name] = sorted(Path(d).glob(f"{name}-*.png"))
            self.assertEqual(len(outs["hs"]), len(outs["chrome"]))
            for a, b in zip(outs["hs"], outs["chrome"]):
                ia, ib = Image.open(a).convert("L"), Image.open(b).convert("L")
                self.assertEqual(ia.size, ib.size)
                diff = ImageChops.difference(ia, ib)
                changed = sum(1 for v in diff.getdata() if v > 40) / (ia.size[0] * ia.size[1])
                self.assertLess(changed, 0.005, f"{a.name} が Chrome 本体と {changed:.2%} 違う")


class CountTitles(unittest.TestCase):
    """check_deck：数で束ねたタイトル（「3つの」「3本の」「3点」）が3枚以上なら FAIL（§4.103）"""

    def test_three_count_titles_fail(self):
        html = GOOD
        for old, new in (("差し戻しの削減と承認の自動化で、承認待ちを当日に縮める", "承認の遅れは3つの壁で起きている"),):
            html = html.replace(old, new, 1)
        html = html.replace("</body>", "".join(f'<section class="s"><h1>{t}</h1><div class="foot"><span>9</span></div></section>'
                                                for t in ("来期は3本柱で進める", "論点は2点に絞られる")) + "</body>", 1)
        self.assertTrue(any("数で束ねたタイトル" in l for l in run(html)[1]))

    def test_ordinal_and_soft_counts_do_not_fail(self):
        # 「第2段階」は序数、「3点を承認いただく」「2段階で進める」は数を材料に言っているだけのこともある → FAIL にしない（WARN 止まり）
        html = inject("差し戻しの削減と承認の自動化で、承認待ちを当日に縮める", "第2段階で全社展開に進む")
        html = html.replace("</body>", "".join(f'<section class="s"><h1>{t}</h1><div class="foot"><span>9</span></div></section>'
                                                for t in ("本日は3点を承認いただく", "導入は2段階で進める", "承認の遅れは3つの壁で起きている")) + "</body>", 1)
        self.assertFalse(any("数で束ねたタイトル" in l for l in run(html)[1]))
        self.assertTrue(any("数を前に出したタイトル" in l for l in warns(html)))

    def test_good_deck_passes(self):
        self.assertFalse(any("数で束ねたタイトル" in l for l in run(GOOD)[1] + warns(GOOD)))


class PartVariant(unittest.TestCase):
    """new_deck.py：b09:axb で軸ボックスの表になる"""

    def test_axb_variant(self):
        r = subprocess.run([sys.executable, str(ROOT / "scripts" / "new_deck.py"), "--show", "b09:axb"], capture_output=True, text=True)
        self.assertIn('<table class="axb"', r.stdout)

    def test_unknown_variant_fails(self):
        r = subprocess.run([sys.executable, str(ROOT / "scripts" / "new_deck.py"), "--show", "b09:zzz"], capture_output=True, text=True)
        self.assertNotEqual(r.returncode, 0)


@unittest.skipUnless(_has_playwright(), "npm run setup で playwright を入れると実行される")
class SpreadItems(LayoutFeedback20261012):
    """check_layout：項目の間を広げて版面を埋める間延び（FAIL）と、縦の線の列が伸びて中が空く（WARN）（§4.12）"""

    def ul(self, spread):
        st = "height:120mm;display:flex;flex-direction:column;justify-content:space-around" if spread else ""
        return f'<ul style="{st};font-size:11pt">' + "".join(f"<li>項目{i}の説明を一文で書く</li>" for i in range(1, 4)) + "</ul>"

    def test_spread_list_fails(self):
        self.assertTrue(any("並べた項目の間" in l and l.startswith("FAIL") for l in self.layout(self.ul(True))))

    def test_compact_list_passes(self):
        self.assertFalse(any("並べた項目の間" in l for l in self.layout(self.ul(False))))

    def cols(self, stretch):
        h = "height:110mm;" if stretch else ""
        col = '<div class="ph" style="{h}border-right:1px solid #ccc;padding:3mm"><b>時期{i}</b><div>中身{i}を書く</div></div>'
        return '<div style="display:grid;grid-template-columns:repeat(4,1fr)">' + "".join(col.format(h=h, i=i) for i in range(1, 5)) + "</div>"

    def test_stretched_columns_warn(self):
        self.assertTrue(any("縦の線で区切った列が伸び" in l for l in self.layout(self.cols(True))))

    def test_content_height_columns_pass(self):
        self.assertFalse(any("縦の線で区切った列が伸び" in l for l in self.layout(self.cols(False))))

    def cards(self, split):
        st = "height:110mm;display:flex;flex-direction:column;" if split else ""
        ul = '<ul style="margin-top:auto">' if split else "<ul>"
        card = f'<div class="cd" style="{st}background:#eee;padding:4mm"><b>柱{{i}}の名前</b><div>柱{{i}}の中身を一文で書く</div>{ul}<li>対象{{i}}</li><li>時期{{i}}</li></ul></div>'
        return '<div style="display:grid;grid-template-columns:repeat(3,1fr);gap:5mm">' + "".join(card.format(i=i) for i in range(1, 4)) + "</div>"

    def test_split_card_fails(self):
        self.assertTrue(any("真ん中が" in l and l.startswith("FAIL") for l in self.layout(self.cards(True))))

    def test_compact_card_passes(self):
        self.assertFalse(any("真ん中が" in l for l in self.layout(self.cards(False))))


@unittest.skipUnless(_has_playwright(), "npm run setup で playwright を入れると実行される")
class TitleSizeAndPdfPages(LayoutFeedback20261012):
    """check_layout：タイトルが22ptから±10%超ずれる（FAIL。§2.1）、PDF のページ数がスライドの枚数と合わない（FAIL）"""

    def full(self, html):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "deck.html"
            p.write_text(html, encoding="utf8")
            r = subprocess.run(["node", str(ROOT / "scripts" / "check_layout.mjs"), str(p)], capture_output=True, text=True, cwd=ROOT, encoding="utf8")
        return r.stdout

    def test_big_title_fails(self):
        html = GOOD.replace("</style>", ".s h1{font-size:30pt}</style>", 1)
        self.assertIn("タイトルの大きさが既定の22pt", self.full(html))

    def test_overflowing_page_splits_pdf(self):
        html = GOOD.replace("</style>", ".s.cover{height:auto;min-height:200mm;overflow:visible}</style>", 1)
        self.assertIn("PDF が", self.full(html))

    def test_source_line_inside_footer_fails(self):
        html = GOOD.replace('<div class="src">出典：架空データ（説明用）</div>', '<div class="src" style="position:absolute;left:16mm;bottom:6mm">出典：架空データ（説明用）</div>', 1)
        self.assertIn("フッターと本文が重なる", self.full(html))

    def test_good_deck_passes(self):
        out = self.full(GOOD)
        self.assertNotIn("タイトルの大きさが既定の22pt", out)
        self.assertNotIn("PDF が", out)


@unittest.skipUnless(_has_playwright(), "npm run setup で playwright を入れると実行される")
class DividerAboveHeader(LayoutFeedback20261012):
    """check_layout：表の列の区切り線・行の軸の縦線が見出しの行まで伸びている（WARN。§6）"""

    def table(self, head_divider):
        th = "border-bottom:2px solid #333;" + ("border-left:1px dotted #999;" if head_divider else "")
        rows = "".join(f'<tr><td>部門{i}</td><td style="border-left:1px dotted #999">月に約{i}0件</td><td style="border-left:1px dotted #999">担当{i}名</td></tr>' for i in range(1, 4))
        return (f'<table style="border-collapse:collapse"><tr><th style="border-bottom:2px solid #333">部門</th><th style="{th}">件数</th>'
                f'<th style="{th}">担当</th></tr>{rows}</table><div style="height:250px;border:1px solid #333">図</div>')

    def test_divider_in_header_warns(self):
        self.assertTrue(any("縦の線が見出しの行まで" in l for l in self.layout(self.table(True))))

    def test_divider_only_in_body_passes(self):
        self.assertFalse(any("縦の線が見出しの行まで" in l for l in self.layout(self.table(False))))

    def axis_table(self, head_line):
        th0 = "border-bottom:2px solid #333;" + ("border-right:2px solid #333;" if head_line else "")
        rows = "".join(f'<tr><td style="border-right:2px solid #333;font-weight:700">部門{i}</td><td>月に約{i}0件</td><td>担当{i}名</td></tr>' for i in range(1, 4))
        return (f'<table style="border-collapse:collapse"><tr><th style="{th0}">部門</th><th style="border-bottom:2px solid #333">件数</th>'
                f'<th style="border-bottom:2px solid #333">担当</th></tr>{rows}</table><div style="height:250px;border:1px solid #333">図</div>')

    def test_axis_line_in_header_warns(self):
        self.assertTrue(any("縦の線が見出しの行まで" in l for l in self.layout(self.axis_table(True))))

    def test_axis_line_only_in_body_passes(self):
        self.assertFalse(any("縦の線が見出しの行まで" in l for l in self.layout(self.axis_table(False))))


class ReviewFixes20261022(unittest.TestCase):
    """check_deck・check_all：中黒1つでタイトルの判定が逆転しない・罫線の階層表記・定型ページの出典照合・sources.md の探し方"""

    T = "差し戻しの削減と承認の自動化で、承認待ちを当日に縮める"

    def page(self, body, cls="s"):
        return f'<section class="{cls}"><h1>検査の見本として置いたページのタイトル</h1>{body}<div class="foot"><span>5</span></div></section>'

    def test_action_title_with_nakaguro_is_warn_not_fail(self):
        # 中黒があっても打ち手の言い切りは FAIL にしない（中黒の有無で判定が逆転しない）
        for t in ("申請・承認の重複した窓口をまとめる", "部門ごとに申請・承認の窓口をまとめる"):
            html = inject(self.T, t)
            self.assertFalse(any("作業の説明" in l for l in run(html)[1]), t)
            self.assertTrue(any("で終わっている" in l for l in warns(html)), t)

    def test_tree_branch_is_not_dash_link(self):
        # 「├─ 項目」のような階層表記は語句のつなぎではない
        html = GOOD.replace("差し戻しの6割は添付漏れ", "差し戻しの6割は添付漏れ<p>├─ 調達コスト</p><p>└─ 物流コスト</p>", 1)
        self.assertFalse(any("「—」で語句をつないでいる" in l for l in run(html)[1]))

    def test_box_drawing_dash_with_spaces_still_fails(self):
        html = GOOD.replace("差し戻しの6割は添付漏れ", "差し戻しの6割は添付漏れ<p>前年比 ─ 横ばい</p>", 1)
        self.assertTrue(any("「—」で語句をつないでいる" in l for l in run(html)[1]))

    def test_sources_skip_std_page(self):
        with tempfile.TemporaryDirectory() as d:
            src = Path(d) / "sources.md"
            src.write_text("# 出典\n- 月約1,200件 社内集計\n", encoding="utf8")
            deck = Path(d) / "deck.html"
            out = []
            for cls in ("s", "s std"):
                deck.write_text("<html><body>" + self.page("<p>申請は月約1,200件</p>") + self.page("<p>従業員は4,321人</p>", cls) + "</body></html>", encoding="utf8")
                r = subprocess.run([sys.executable, str(CHECK), str(deck), *NO_BRAND, "--sources", str(src)], capture_output=True, text=True)
                out.append(r.stdout)
        self.assertIn("「4,321人」", out[0])
        self.assertNotIn("「4,321人」", out[1])

    def test_cover_title_and_subtitle_do_not_warn(self):
        # class に cover/chap/back が付いたページのタイトル無し・サブタイトルは WARN にしない。本文ページのサブタイトルは WARN
        w = warns(GOOD)
        self.assertFalse(any("タイトルが空" in l or "`.sub`" in l for l in w), w)
        w = warns(GOOD.replace("</body>", self.page('<div class="sub">本文ページの補足</div>') + "</body>", 1))
        self.assertTrue(any("`.sub` ×1" in l for l in w), w)

    def test_check_all_sources_without_value_exits_with_reason(self):
        with tempfile.TemporaryDirectory() as d:
            deck = Path(d) / "deck.html"
            deck.write_text(GOOD, encoding="utf8")
            r = subprocess.run([sys.executable, str(ROOT / "scripts" / "check_all.py"), str(deck), "--sources"],
                               capture_output=True, text=True)
        self.assertNotIn("Traceback", r.stderr)
        self.assertIn("--sources のあとに", r.stderr)
        self.assertNotEqual(r.returncode, 0)

    @unittest.skipUnless(_has_playwright(), "playwright が要る")
    def test_check_all_ignores_sources_in_cwd(self):
        # 作業フォルダの sources.md（別のデッキの控えかもしれない）では照合しない
        with tempfile.TemporaryDirectory() as d, tempfile.TemporaryDirectory() as cwd:
            deck = Path(d) / "deck.html"
            deck.write_text(GOOD, encoding="utf8")
            (Path(cwd) / "sources.md").write_text("# 出典\n- 別のデッキの数字 99件\n", encoding="utf8")
            r = subprocess.run([sys.executable, str(ROOT / "scripts" / "check_all.py"), str(deck)], capture_output=True, text=True, cwd=cwd)
        self.assertIn("出典の照合なし", r.stdout)
        self.assertNotIn("出典の控えに無い数字", r.stdout)

    @unittest.skipUnless(_has_playwright(), "playwright が要る")
    def test_good_deck_has_no_warn(self):
        # 良い見本は check_all で FAIL も WARN も出ない
        with tempfile.TemporaryDirectory() as d:
            deck = Path(d) / "deck.html"
            deck.write_text(GOOD, encoding="utf8")
            r = subprocess.run([sys.executable, str(ROOT / "scripts" / "check_all.py"), str(deck)], capture_output=True, text=True, cwd=d)
        self.assertIn("0 FAIL / 0 WARN", r.stdout, r.stdout[-1500:])

    @unittest.skipUnless(_has_playwright(), "playwright が要る")
    def test_check_all_lists_pages_beyond_twenty(self):
        # check_layout の表示上限（各20件）で21件目以降のページが一覧から消えない
        pages = "".join(self.page('<div class="c"><div style="height:320px;border:1px solid #333">図</div></div>') for _ in range(24))
        with tempfile.TemporaryDirectory() as d:
            deck = Path(d) / "deck.html"
            deck.write_text(GOOD.replace("</body>", pages + "</body>", 1), encoding="utf8")
            r = subprocess.run([sys.executable, str(ROOT / "scripts" / "check_all.py"), str(deck)], capture_output=True, text=True, cwd=d)
        self.assertIn("== p30", r.stdout)
        self.assertRegex(r.stdout, r"== p30\n(  (FAIL|WARN) .*\n)*  FAIL  版面の")


@unittest.skipUnless(_has_playwright(), "npm run setup で playwright を入れると実行される")
class LayoutReviewFixes20261022(LayoutFeedback20261012):
    """check_layout：zoom の試算（本文の大きさの揃い・目標30%・拡大前からのはみ出し）・2行にまたがるセルと罫線"""

    test_right_side_gap_warns = test_two_highlighted_rows_warn = test_std_page_skips_layout_checks = None
    test_unequal_contrast_columns_warn = test_equal_contrast_columns_pass = None

    def hint(self, body):
        return next((l for l in self.layout(body) if "版面の" in l), "")

    def test_zoom_hint_keeps_body_size_in_range(self):
        # デッキの中央値より少し大きい本文のページで、勧めた倍率に従うと本文の大きさの WARN が出ない
        ul = ('<ul style="font-size:19px;margin:0"><li>申請の内容を確認して承認者に回す</li>'
              '<li>添付漏れを提出前に警告して直させる</li><li>少額の申請は条件を満たせば自動で承認する</li></ul>')
        line = self.hint(f'<div class="c">{ul}</div>')
        m = re.search(r"zoom:([\d.]+)|上限の([\d.]+)倍", line)
        self.assertTrue(m, line)
        z = m.group(1) or m.group(2)
        out = self.layout(f'<div class="c" style="zoom:{z}">{ul}</div>')
        self.assertFalse(any("本文の実際の大きさ" in l for l in out), (z, out))

    def test_zoom_hint_targets_warn_threshold(self):
        # 「直る」と案内するのは空きが WARN の閾値（30%）を切るときだけ
        line = self.hint('<div class="c"><div style="height:320px;background:#eee"></div></div>')
        self.assertTrue(line)
        self.assertFalse(re.search(r"を付けると空き約(3[1-9]|[4-9]\d)%", line), line)

    def test_zoom_hint_ignores_overflow_before_zoom(self):
        # 拡大前からフッターに入り込んでいる文字（別の FAIL で出る）で「1.05倍で即はみ出す」と案内しない
        body = ('<div class="c"><div style="height:200px;background:#eee"></div></div>'
                '<div style="position:absolute;left:120mm;bottom:6mm;font-size:6pt;line-height:1">注記の文字</div>')
        line = self.hint(body)
        self.assertTrue(line)
        self.assertNotIn("1.05倍にするだけで版面からはみ出す", line)

    ROWSPAN = ('<style>.rs tr+tr{{border-top:1px solid #333}}</style><div class="c"><table class="rs" style="font-size:13pt">'
               '<tr><td {a} style="width:60mm">全部門に共通する条件</td><td>営業部の申請を確認する</td><td>月に約120件</td></tr>'
               '<tr>{b}<td>本社の申請を確認する</td><td>月に約300件</td></tr>'
               '<tr><td>拠点ごとの条件</td><td>全拠点の申請を確認する</td><td>月に約500件</td></tr></table></div>')

    def test_rowspan_cell_is_not_crossed_by_row_rule(self):
        out = self.layout(self.ROWSPAN.format(a='rowspan="2"', b=""))
        self.assertFalse(any("罫線に重なっている" in l for l in out), out)

    def test_text_over_row_rule_still_warns(self):
        # またがるセルでない文字が罫線に重なるときは、これまでどおり WARN
        body = ('<div style="height:40px;border-bottom:2px solid #333">指標名、単位</div>'
                '<div style="margin-top:-18px;font-size:20px;font-weight:700">48.5</div><div style="height:300px;border:1px solid #333">図</div>')
        self.assertTrue(any("罫線に重なっている" in l for l in self.layout(body)))


class AxisNameInHeaderBand(unittest.TestCase):
    # §4.102: 見出し行の左端セルを空欄にしない
    def test_empty_header_corner_fires(self):
        code, fails = run(inject('<th class="ax" style="width:52mm">段階</th>', '<th class="ax" style="width:52mm"></th>'))
        self.assertEqual(code, 1)
        self.assertTrue(any("§4.102" in f for f in fails), fails)

    def test_named_header_corner_passes(self):
        code, fails = run(GOOD)
        self.assertFalse(any("§4.102" in f for f in fails), fails)


class AppendixDivider(unittest.TestCase):
    # §4.100: 本編と付録の間に区切りのスライドを挟む
    APPX = '<section class="s"><div class="tag">付録</div><h1>付録：試行した営業部では、承認待ちの平均が3日から当日に縮んだ</h1><p>営業部は平均3日から当日に縮んだ</p><div class="foot"><b>業務自動化の検討</b><span>5</span></div></section>\n'
    DIV = '<section class="s"><h1>Appendix</h1><p>部門別の数字</p></section>\n'
    TABLE = '<table><tr><th>部門</th><th>試行前</th><th>試行後</th></tr><tr><td>営業部</td><td>3日</td><td>当日</td></tr></table>'

    def _deck(self, extra):
        # 裏表紙（最後の cover）の手前に差し込む
        head, sep, tail = GOOD.rpartition('<section class="s cover">')
        return head + extra + sep + tail

    def _fires(self, extra):
        code, fails = run(self._deck(extra))
        return any("§4.100" in f for f in fails), fails

    def test_appendix_without_divider_fires(self):
        fired, fails = self._fires(self.APPX)
        self.assertTrue(fired, fails)

    def test_appendix_with_divider_passes(self):
        fired, fails = self._fires(self.DIV + self.APPX)
        self.assertFalse(fired, fails)

    def test_labelled_dividers_pass(self):
        # 区切りのタイトルに番号・短い名詞句が付いていても区切りとして数える
        for ttl in ("Appendix A", "Appendix 1", "付録 A", "Appendix｜部門別の数字", "APPENDIX：参考の数字"):
            with self.subTest(ttl=ttl):
                div = f'<section class="s"><h1>{ttl}</h1><p>部門別の数字</p></section>\n'
                fired, fails = self._fires(div + self.APPX)
                self.assertFalse(fired, fails)

    def test_short_appendix_page_is_not_its_own_divider(self):
        # 「付録：」の後ろが短い付録ページは、区切りが無ければ FAIL（自分自身を区切りに数えない）
        page = '<section class="s"><h1>付録：営業部では承認待ちが当日に縮んだ</h1><p>営業部は平均3日から当日に縮んだ</p><div class="foot"><b>業務自動化の検討</b><span>5</span></div></section>\n'
        fired, fails = self._fires(page)
        self.assertTrue(fired, fails)

    def test_divider_with_table_is_not_a_divider(self):
        # タイトルが「Appendix」でも表が載っていれば付録の中身のページ
        page = f'<section class="s"><h1>Appendix</h1>{self.TABLE}<div class="foot"><b>業務自動化の検討</b><span>5</span></div></section>\n'
        fired, fails = self._fires(page)
        self.assertTrue(fired, fails)

    def test_main_body_title_starting_with_furoku_passes(self):
        # 本編で「付録の手順書…」のように付録で始まるタイトルは付録ページではない
        page = '<section class="s"><h1>付録の手順書を改め、承認の待ち時間を当日に縮める</h1><p>営業部から順に手順書を差し替える</p><div class="foot"><b>業務自動化の検討</b><span>5</span></div></section>\n'
        fired, fails = self._fires(page)
        self.assertFalse(fired, fails)

    def test_tag_chip_after_bar_fires(self):
        # 資料名・日付の帯（.bar）から始まるページでも、タグチップ「appendix」（小文字）で付録と分かる
        page = ('<section class="s"><div class="bar"><div class="logo">資料名</div><div class="date">2026.10</div></div>'
                f'<div class="tag">appendix</div><h1>試行した営業部では、承認待ちの平均が3日から当日に縮んだ</h1>{self.TABLE}'
                '<div class="foot"><b>業務自動化の検討</b><span>5</span></div></section>\n')
        fired, fails = self._fires(page)
        self.assertTrue(fired, fails)
        fired, fails = self._fires(self.DIV + page)
        self.assertFalse(fired, fails)

    def test_agenda_listing_appendix_is_not_an_appendix_page(self):
        # アジェンダの項目・表のセルの「Appendix」「付録」はタグではない（後ろに正しい区切りがあれば通る）
        agendas = (
            '<section class="s"><h1>アジェンダ</h1><ol><li>現状の課題</li><li>打ち手</li><li>Appendix</li></ol></section>\n',
            '<section class="s"><h1>営業部から順に承認の手順を差し替える</h1><table><tr><th>章</th><th>内容</th></tr>'
            '<tr><td>付録</td><td>部門別の数字</td></tr></table><div class="foot"><b>業務自動化の検討</b><span>3</span></div></section>\n',
        )
        for agenda in agendas:
            with self.subTest(agenda=agenda[:40]):
                head, sep, tail = GOOD.partition('</section>')
                deck = head + sep + "\n" + agenda + tail
                head, sep, tail = deck.rpartition('<section class="s cover">')
                code, fails = run(head + self.DIV + self.APPX + sep + tail)
                self.assertFalse(any("§4.100" in f for f in fails), fails)

    def test_divider_with_bracket_passes(self):
        # 「参考資料（部門別）」のように括弧が続く区切りも区切りとして数える
        div = '<section class="s"><h1>参考資料（部門別）</h1><p>部門別の数字</p></section>\n'
        fired, fails = self._fires(div + self.APPX)
        self.assertFalse(fired, fails)


def run_raw(html, files, *extra):
    """run_files と同じだが、終了コード・標準出力・標準エラーをそのまま返す"""
    with tempfile.TemporaryDirectory() as d:
        for name, body in files.items():
            (Path(d) / name).write_text(body, encoding="utf8")
        p = Path(d) / "deck.html"
        p.write_text(html, encoding="utf8")
        extra = [str(Path(d) / a[1:]) if a.startswith("@") else a for a in extra]
        return subprocess.run([sys.executable, str(CHECK), str(p), *extra], capture_output=True, text=True)


class OldColors(unittest.TestCase):
    # 自社のレガシー色は本体に書かず、local/old-colors.txt か --old-colors で渡す
    def test_listed_color_fires(self):
        code, fails = run_files(GOOD, {"old.colors": "# 旧ブランド色\n#8A8F94\n"}, "--old-colors", "@old.colors")
        self.assertEqual(code, 1)
        self.assertTrue(any("8A8F94" in f for f in fails), fails)

    def test_brand_none_skips(self):
        code, fails = run_files(GOOD, {"old.colors": "8A8F94\n"}, "--old-colors", "@old.colors", "--brand", "none")
        self.assertFalse(any("禁止色" in f for f in fails), fails)

    def test_no_list_no_check(self):
        code, fails = run(GOOD)
        self.assertFalse(any("禁止色" in f for f in fails), fails)

    def test_line_formats_are_read(self):
        # 後ろの説明・空白1つの「# コメント」・BOM 付き・小文字でも色を読む
        for body in ("8A8F94 # old\n", "﻿8A8F94\n", "#8A8F94 旧グレー\n", "8a8f94\n", "  #8A8F94\t# 旧\n"):
            with self.subTest(body=body):
                code, fails = run_files(GOOD, {"old.colors": body}, "--old-colors", "@old.colors")
                self.assertTrue(any("8A8F94" in f for f in fails), fails)

    def test_unreadable_line_warns_and_short_value_does_not_match(self):
        # 6桁でない値は読み飛ばして WARN。短い値で部分一致させない。「#」付きの3桁もコメント扱いにしない
        for body in ("8A8\n", "#8A8 旧色\n"):
            with self.subTest(body=body):
                r = run_raw(GOOD, {"old.colors": body}, "--old-colors", "@old.colors")
                self.assertFalse(any("禁止色" in l for l in r.stdout.splitlines() if l.startswith("FAIL")), r.stdout[-400:])
                self.assertTrue(any(l.startswith("WARN") and "読めない" in l for l in r.stdout.splitlines()), r.stdout[-400:])

    def test_eight_digit_hex_matches(self):
        # 透明度つきの8桁（デッキ側・一覧側とも）は先頭6桁の色として扱う
        html = GOOD.replace("#8a8f94", "#8a8f94cc")
        for body in ("8A8F94\n", "#8A8F94CC\n"):
            with self.subTest(body=body):
                code, fails = run_files(html, {"old.colors": body}, "--old-colors", "@old.colors")
                self.assertTrue(any("8A8F94" in f for f in fails), fails)

    def test_brand_none_with_old_colors_says_so(self):
        r = run_raw(GOOD, {"old.colors": "8A8F94\n"}, "--old-colors", "@old.colors", "--brand", "none")
        self.assertIn("--brand none", r.stderr)

    def test_color_inside_longer_hex_does_not_match(self):
        html = GOOD.replace("#8a8f94", "#8a8f9")  # 8A8F94 を消す
        html = html.replace("</style>", ".z{color:#18A8F94}</style>", 1)  # 前後に16進が続く並び
        code, fails = run_files(html, {"old.colors": "8A8F94\n"}, "--old-colors", "@old.colors")
        self.assertFalse(any("禁止色" in f for f in fails), fails)


class CheckDeckArgs(unittest.TestCase):
    # 引数・設定ファイルの誤りは traceback を出さず、理由を出して終える
    def _assert_usage_error(self, r, word):
        self.assertEqual(r.returncode, 2, r.stdout[-300:] + r.stderr[-300:])
        self.assertNotIn("Traceback", r.stderr)
        self.assertIn(word, r.stderr)

    def test_missing_old_colors_file(self):
        self._assert_usage_error(run_raw(GOOD, {}, "--old-colors", "@nothing.txt"), "見つからない")

    def test_brand_without_value(self):
        self._assert_usage_error(run_raw(GOOD, {}, "--brand"), "--brand")

    def test_unknown_brand_value(self):
        self._assert_usage_error(run_raw(GOOD, {}, "--brand", "acme"), "acme")


@unittest.skipUnless(_has_pptx(), "python-pptx が入っていると実行される")
class PptxAppendixAndAxisNames(unittest.TestCase):
    """§4.100・§4.102 の PPTX 版"""

    def _check(self, pages):
        """pages: [(タイトル, 本文 or None, 表 or None)]。表は (見出しを太字にするか, firstRow, 行のリスト)"""
        from pptx import Presentation
        from pptx.util import Inches, Pt
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        for title, body, table, *tag in pages:
            s = prs.slides.add_slide(prs.slide_layouts[5])
            s.shapes.title.text = title
            if tag:  # 上端の小さいタグチップ
                s.shapes.add_textbox(Inches(0.5), Inches(0.15), Inches(1.2), Inches(0.4)).text_frame.text = tag[0]
            if body:
                s.shapes.add_textbox(Inches(1), Inches(2), Inches(10), Inches(1)).text_frame.text = body
            if table:
                bold_head, first_row, rows = table
                gf = s.shapes.add_table(len(rows), len(rows[0]), Inches(1), Inches(3), Inches(10), Inches(2))
                gf.table.first_row = first_row
                for r, row in enumerate(rows):
                    for c, v in enumerate(row):
                        cell = gf.table.cell(r, c)
                        cell.text = v
                        for para in cell.text_frame.paragraphs:
                            for run_ in para.runs:
                                run_.font.bold = bold_head and r == 0
                                run_.font.size = Pt(12)
        with tempfile.TemporaryDirectory() as d:
            path = Path(d) / "deck.pptx"
            prs.save(path)
            r = subprocess.run([sys.executable, str(CHECK), str(path), *NO_BRAND], capture_output=True, text=True)
        return [l for l in r.stdout.splitlines() if l.startswith("FAIL")]

    BODY = ("営業部の承認待ちを当日に縮める", "営業部から順に差し替える", None)
    ROWS = [["部門", "試行前", "試行後"], ["営業部", "3日", "当日"]]
    NO_CORNER = [["", "試行前", "試行後"], ["営業部", "3日", "当日"]]

    def _deck(self, *pages):
        return [("業務自動化の検討", "2026年10月", None), self.BODY, *pages]

    def test_appendix_without_divider_fires(self):
        fails = self._check(self._deck(("付録：試行した営業部では承認待ちの平均が当日に縮んだ", None, (True, False, self.ROWS))))
        self.assertTrue(any("§4.100" in f for f in fails), fails)

    def test_labelled_divider_passes(self):
        for ttl in ("Appendix A", "付録 A", "Appendix｜部門別の数字"):
            with self.subTest(ttl=ttl):
                fails = self._check(self._deck((ttl, "部門別の数字", None),
                                               ("付録：試行した営業部では承認待ちの平均が当日に縮んだ", None, (True, False, self.ROWS))))
                self.assertFalse(any("§4.100" in f for f in fails), fails)

    def test_agenda_item_is_not_a_tag(self):
        # アジェンダや本文にある「Appendix」だけの図形はタグではない（後ろに区切りがあれば通る）
        appx = ("付録：試行した営業部では承認待ちの平均が当日に縮んだ", None, (True, False, self.ROWS))
        for page in (("アジェンダ", "Appendix", None), ("営業部から順に承認の手順を差し替える", "Appendix", None)):
            with self.subTest(title=page[0]):
                fails = self._check(self._deck(page, ("Appendix", "部門別の数字", None), appx))
                self.assertFalse(any("§4.100" in f for f in fails), fails)

    def test_tag_chip_at_top_fires(self):
        page = ("試行した営業部では承認待ちの平均が当日に縮んだ", None, (True, False, self.ROWS), "付録")
        fails = self._check(self._deck(page))
        self.assertTrue(any("§4.100" in f for f in fails), fails)
        fails = self._check(self._deck(("Appendix", "部門別の数字", None), page))
        self.assertFalse(any("§4.100" in f for f in fails), fails)

    def test_short_appendix_page_fires(self):
        fails = self._check(self._deck(("付録：営業部では承認待ちが当日に縮んだ", "営業部は平均3日から当日に縮んだ", None)))
        self.assertTrue(any("§4.100" in f for f in fails), fails)

    def test_main_body_title_starting_with_furoku_passes(self):
        fails = self._check(self._deck(("付録の手順書を改め、承認の待ち時間を当日に縮める", "営業部から順に手順書を差し替える", None)))
        self.assertFalse(any("§4.100" in f for f in fails), fails)

    def test_empty_corner_in_header_fires(self):
        for first_row, bold in ((False, True), (True, False)):
            with self.subTest(first_row=first_row, bold=bold):
                fails = self._check(self._deck(("部門ごとに承認待ちの日数を比べる", None, (bold, first_row, self.NO_CORNER))))
                self.assertTrue(any("§4.102" in f for f in fails), fails)

    def test_table_without_header_passes(self):
        # 見出しの無い表（太字でも firstRow でもない）は、左上が空欄でも対象外
        fails = self._check(self._deck(("部門ごとに承認待ちの日数を比べる", None, (False, False, self.NO_CORNER))))
        self.assertFalse(any("§4.102" in f for f in fails), fails)

    def test_named_corner_passes(self):
        fails = self._check(self._deck(("部門ごとに承認待ちの日数を比べる", None, (True, True, self.ROWS))))
        self.assertFalse(any("§4.102" in f for f in fails), fails)
